"""WS-6 · Informes Excel con datos del flujo (copia QA, escribe datos propios vía la app). Incluye CR-10 (correctivo).

Para L-6 (unidad y laboratorio), L-7, L-8, L-3 y Adquisición/L-12, en un plan PREVENTIVO completado y uno CORRECTIVO completado:
  IN-01 200 + xlsx que abre con openpyxl y nombre de archivo en Content-Disposition
  IN-02 contiene los datos clave del registro (SQL): inventario, texto de la falla / checklist, costo real
  IN-03 el L-7 trae el checklist técnico completo (13 puntos)
  IN-07 impresión: área de impresión, orientación y ajuste a página iguales a la plantilla de wwwroot/templates
  CR-10 correctivo: L-6 → 400 «solo aplica al proceso preventivo»; los demás muestran el proceso «correctivo»
  IN-04 ids inexistentes → 400/404 con mensaje, nunca 500
  IN-05 permisos: Supervisor descarga; Cierre de práctica y anónimo no
"""
import io
import os
import urllib.error
import urllib.request
from decimal import Decimal
from pathlib import Path

import openpyxl

import e2e_lib
import http_cases_correctivo_common as C
import http_cases_procesos_common as P
from http_cases_procesos_common import check, plan, scalar

TEMPLATES = Path(__file__).resolve().parents[2] / "wwwroot" / "templates"
PLANTILLA = {"l6": "L6V2.xlsx",   # ReportService.L6 usa L6V2.xlsx
              "l7": "L7.xlsx", "l8": "L8.xlsx", "l3": "L3.xlsx", "adquisicion": "Adquisicion.xlsx"}


def fetch(opener, path):
    try:
        r = opener.open(e2e_lib.BASE + path, timeout=180)
        return r.status, r.headers.get("Content-Type", ""), r.headers.get("Content-Disposition", ""), r.read(), r.geturl()
    except urllib.error.HTTPError as e:
        return e.code, e.headers.get("Content-Type", ""), "", e.read(), path


def libro(body):
    try:
        return openpyxl.load_workbook(io.BytesIO(body), data_only=True)
    except Exception:
        return None


def texto(wb):
    return " ".join(str(v) for ws in wb.worksheets for row in ws.iter_rows(values_only=True) for v in row if v is not None)


def numeros(wb):
    return [v for ws in wb.worksheets for row in ws.iter_rows(values_only=True) for v in row if isinstance(v, (int, float))]


def impresion(ws):
    ps = ws.page_setup
    fit = bool(ws.sheet_properties.pageSetUpPr and ws.sheet_properties.pageSetUpPr.fitToPage)
    return (str(ws.print_area or ""), ps.orientation, fit, ps.fitToWidth if fit else None, ps.fitToHeight if fit else None)


def anon_status(path):
    class NoRedir(urllib.request.HTTPRedirectHandler):
        def redirect_request(self, *a, **k):
            return None
    op = urllib.request.build_opener(NoRedir)
    try:
        r = op.open(e2e_lib.BASE + path, timeout=60)
        return r.status, r.headers.get("Location", "")
    except urllib.error.HTTPError as e:
        return e.code, e.headers.get("Location", "")


def informes_de(pid, mgmt, lab, unit):
    p = plan(pid)
    return {
        "l6": f"/api/reports/download/l6?managementId={mgmt}&labId={lab}&unitId={unit}",
        "l7": f"/api/reports/download/l7?requestId={p['req']}&managementId={mgmt}",
        "l8": f"/api/reports/download/l8?managementPlanId={pid}",
        "l3": f"/api/reports/download/l3?departureId={p['dep']}",
        "adquisicion": f"/api/reports/download/adquisicion?requestId={p['acq']}",
    }


def run(R, c):
    lab = P.lab_propio(c, "informes")
    corr = C.asegurar_correctiva(c)
    prev = P.llevar_a(c, lab, "inp", "P7")
    cor = C.llevar_corr(c, corr, lab, "inc", "DONE")
    casos = (("preventivo", prev, P.MGMT), ("correctivo", cor, corr))
    plantillas = {k: openpyxl.load_workbook(TEMPLATES / v).worksheets[0] for k, v in PLANTILLA.items()}

    for proceso, ctx, mgmt in casos:
        pid, unit = ctx["plan"], ctx["unit"]
        p = plan(pid)
        inv = scalar("SELECT InventoryNumber FROM EquipmentUnits WHERE Id = ?", unit)
        costo = scalar("SELECT ActualCost FROM Maintenances WHERE Id = ?", p["mnt"])
        falla = scalar("SELECT Description FROM Requests WHERE Id = ?", p["req"])
        for clave, path in informes_de(pid, mgmt, lab, unit).items():
            tag = f"[{proceso} {clave.upper()}]"
            st, ct, cd, body, _ = fetch(c.opener, path)
            if clave == "l6" and proceso == "correctivo":
                check(R, f"CR-10 {tag} L-6 en correctivo → 400 «solo aplica al proceso preventivo»",
                      st == 400 and "preventivo" in body.decode("utf-8", "ignore").lower(), f"status={st} {body[:120]!r}")
                continue
            wb = libro(body) if st == 200 else None
            check(R, f"IN-01 {tag} 200, xlsx válido y nombre de archivo", st == 200 and wb is not None and "spreadsheetml" in ct and "filename" in cd.lower(),
                  f"status={st} ct={ct} cd={cd[:60]} {body[:100]!r}" if wb is None else f"cd={cd[:60]}")
            if wb is None:
                continue
            t = texto(wb)
            if clave == "adquisicion":
                # formulario institucional de compra: lista conceptos de costo, no la unidad
                check(R, f"IN-02 {tag} contiene el concepto de costo del L-8 (Repuesto QAP)", "repuesto qap" in t.lower(), f"texto={t[:160]}")
                check(R, f"IN-03 {tag} sin texto de ejemplo de la plantilla («TERMOTANQUE»)", "termotanque" not in t.lower(), f"texto={t[:200]}")
            else:
                check(R, f"IN-02 {tag} contiene el inventario de la unidad ({inv})", inv.lower() in t.lower(), f"texto={t[:120]}")
            if clave == "l3":
                g = e2e_lib.db_rows("SELECT Year, Semester FROM Managements WHERE Id = ?", mgmt)[0]
                esperado_g = f"{'I' * int(g[1])}/{g[0]}"
                check(R, f"IN-02 {tag} la gestión impresa es la del plan ({esperado_g})", f"GESTIÓN: {esperado_g}" in t or f"GESTION: {esperado_g}" in t,
                      f"texto gestión={[x for x in t.split(' | ') if 'GESTI' in x.upper()][:2] or t[t.upper().find('GESTI'):t.upper().find('GESTI') + 20]}")
            if clave == "l7":
                check(R, f"IN-02 {tag} contiene la descripción de la falla", (falla or "")[:30] in t, f"falla={falla!r}")
                check(R, f"IN-03 {tag} trae el checklist técnico completo (13 puntos)", "13." in t and "1." in t, f"texto={t[:160]}")
            if clave == "l8" and costo is not None:
                check(R, f"IN-02 {tag} contiene el costo real ({costo})",
                      any(abs(Decimal(str(v)) - Decimal(str(costo))) <= Decimal("0.01") for v in numeros(wb)) or f"{Decimal(str(costo)):.2f}" in t,
                      f"numeros={numeros(wb)[:10]}")
            if proceso == "correctivo" and clave in ("l7", "l8"):   # L-3 y Adquisición son formularios institucionales sin campo de proceso
                check(R, f"CR-10 {tag} el informe indica el proceso correctivo", "correctiv" in t.lower(), f"texto={t[:160]}")
            if clave in plantillas:
                esperado, obtenido = impresion(plantillas[clave]), impresion(wb.worksheets[0])
                check(R, f"IN-07 {tag} configuración de impresión igual a la plantilla (área, orientación, ajuste)", esperado == obtenido,
                      f"plantilla={esperado} generado={obtenido}")

    # ---------------- IN-04 ids inexistentes
    for nombre, path in (("L-7 requestId inexistente", "/api/reports/download/l7?requestId=999999"),
                         ("L-8 managementPlanId inexistente", "/api/reports/download/l8?managementPlanId=999999"),
                         ("L-3 departureId inexistente", "/api/reports/download/l3?departureId=999999"),
                         ("Adquisición requestId inexistente", "/api/reports/download/adquisicion?requestId=999999"),
                         ("L-6 laboratorio inexistente", f"/api/reports/download/l6?managementId={P.MGMT}&labId=999999"),
                         ("L-8 sin parámetros", "/api/reports/download/l8")):
        st, _, _, body, _ = fetch(c.opener, path)
        check(R, f"IN-04 {nombre} → 400/404 con mensaje (nunca 500)", st in (400, 404) and len(body) > 0, f"status={st} {body[:100]!r}")

    # ---------------- IN-05 permisos
    path = informes_de(prev["plan"], P.MGMT, lab, prev["unit"])["l8"]
    sup = P.client("qa_sup", os.environ["E2E_PW"])
    st, ct, _, _, _ = fetch(sup.opener, path)
    check(R, "IN-05 Supervisor descarga el L-8", st == 200 and "spreadsheetml" in ct, f"status={st} ct={ct}")
    cierre = P.client("qa_cierre", os.environ["E2E_PW"])
    st, ct, _, _, url = fetch(cierre.opener, path)
    check(R, "IN-05 Cierre de práctica NO descarga el L-8", "spreadsheetml" not in ct, f"status={st} ct={ct} url={url}")
    st, loc = anon_status(path)
    check(R, "IN-05 anónimo NO descarga el L-8 (302 a /Login o 401)", st in (302, 401) and (st == 401 or "login" in loc.lower()), f"status={st} loc={loc}")


if __name__ == "__main__":
    P.run_module("informes", run)
