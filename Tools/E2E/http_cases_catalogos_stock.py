"""
Catálogos / Existencias (Index, Edit con concurrencia y returnUrl, Import con vista previa y aplicar): reglas del servidor por HTTP.
Solo contra la copia QA. Usa artículos QAC propios (creados aquí) para no alterar las existencias reales de la copia.
La subida multipart se implementa en http_cases_catalogos_common.multipart_post (el motor solo envía urlencoded).
"""
import base64
import io
import re
import urllib.parse

import openpyxl
from lxml import html as LH

from http_cases_catalogos_common import (check, check_404, check_anon, expect, get_nofollow, multipart_post, run_module, scalar, sql, tag, tokenless, uniq)
from http_forms import parse_messages

B = "/Catalogs/Stock"
LABEL = "Stock"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
HEAD = ["LaboratorioCodigo", "ArticuloCodigo", "CantidadDisponible"]


def L(s):
    return f"{LABEL}/{s}"


def stock(lab, art):
    r = sql("select Id, AvailableQuantity, RowVersion, UpdatedOn from ArticleStocks where LaboratoryId=? and ArticleId=?", lab, art)
    return tuple(r[0]) if r else None


def xlsx_bytes(rows, header=HEAD, sheet="ExistenciasArticulos"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet
    if header is not None:
        ws.append(header)
    for r in rows:
        ws.append(list(r))
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


def preview_summary(html_text):
    """(nuevas, a actualizar, sin cambios, con error, token, filas[(fila, accion, texto)])."""
    nums = re.findall(r'<div class="h4 font-weight-bold text-(success|warning|secondary|danger) mb-0">(\d+)</div>', html_text)
    d = {k: int(v) for k, v in nums}
    doc = LH.fromstring(html_text)
    rows = [(r.get("data-action"), re.sub(r"\s+", " ", r.text_content()).strip()) for r in doc.xpath("//tr[contains(@class,'si-row')]")]
    tok = re.search(r'name="Token" value="([^"]+)"', html_text)
    return d.get("success"), d.get("warning"), d.get("secondary"), d.get("danger"), tok.group(1) if tok else None, rows


def run(R, c):
    for p in (B, f"{B}/Import", f"{B}/Edit/1/1", f"{B}/Import?handler=Template"):
        check_anon(R, LABEL, p)
    check_anon(R, LABEL, f"{B}/Edit/1/1", post=True)
    check_anon(R, LABEL, f"{B}/Import?handler=Preview", post=True)
    check_anon(R, LABEL, f"{B}/Import?handler=Apply", post=True)

    cat = scalar("select top 1 Id from ArticleCategories where Status=0 order by Id")
    unit = scalar("select top 1 Id from MeasurementUnits where Status=0 order by Id")
    labs = [r[0] for r in sql("select Id from Laboratories where Status=0 order by Id")]
    LAB1, LAB2 = labs[0], labs[1]
    lab1_code = scalar("select Code from Laboratories where Id=?", LAB1)
    lab1_name = scalar("select Name from Laboratories where Id=?", LAB1)
    stock0 = scalar("select count(*) from ArticleStocks")

    arts = []
    for i in range(6):
        nm = uniq(f"QAC Stk {i}")
        c.submit("/Catalogs/Articles/Create", {"Input.Name": nm, "Input.ArticleCategoryId": str(cat), "Input.MeasurementUnitId": str(unit)})
        aid = scalar("select top 1 Id from Articles where Name=? order by Id desc", nm)
        arts.append((aid, nm, scalar("select Code from Articles where Id=?", aid)))
    check(R, L("Preparación: 6 artículos QAC creados"), all(a[0] for a in arts))
    A0, A0N, A0C = arts[0]
    imp = arts[1:]

    # ---------------------------------------------------------------- Index
    st, page = c.get(f"{B}?{urllib.parse.urlencode({'LaboratoryFilter': LAB1, 'SearchTerm': A0N})}")
    check(R, L("Index: artículo nuevo aparece 'sin registro' para el laboratorio"), st == 200 and f"{B}/Edit/{LAB1}/{A0}" in page and "sin registro" in page.lower(), f"status={st}")
    for q in ({"LaboratoryFilter": "abc"}, {"LaboratoryFilter": "99999"}, {"LaboratoryFilter": "0"}, {"LaboratoryFilter": "-1"}, {"SearchTerm": "%'<>"}, {"SearchTerm": "x" * 400}):
        st, _ = c.get(f"{B}?{urllib.parse.urlencode({**q, 'SearchTerm': q.get('SearchTerm', 'QACzzz')})}")
        check(R, L(f"Index: {str(q)[:60]} responde 200"), st == 200, f"status={st}")
    st, page = c.get(f"{B}?SearchTerm=QACzzzz")
    check(R, L("Index sin laboratorio elegido usa ECONO por defecto"), st == 200 and "ECONO" in page, f"status={st}")

    # ---------------------------------------------------------------- Edit: GET
    st, text = c.text(f"{B}/Edit/{LAB1}/{A0}")
    check(R, L("Edit GET fila nueva: 'Registrar Existencia'"), st == 200 and "Registrar Existencia" in text, text[:100])
    action, fields = c.form(f"{B}/Edit/{LAB1}/{A0}")
    print(f"[{LABEL}] campos Edit:", [(k, v[:14]) for k, v in fields if k != "__RequestVerificationToken"])
    for bad in (f"{B}/Edit/abc/1", f"{B}/Edit/1/abc", f"{B}/Edit/1", f"{B}/Edit/-1/-1x"):
        check_404(R, c, LABEL, bad)
    for path, msg in ((f"{B}/Edit/99999/{A0}", "El laboratorio seleccionado no existe o no está activo."), (f"{B}/Edit/0/{A0}", "El laboratorio seleccionado no existe o no está activo."),
                      (f"{B}/Edit/-1/{A0}", "El laboratorio seleccionado no existe o no está activo."), (f"{B}/Edit/{LAB1}/99999999", "El artículo seleccionado no existe o no está activo."),
                      (f"{B}/Edit/{LAB1}/0", "El artículo seleccionado no existe o no está activo.")):
        st, page = c.get(path)
        check(R, L(f"Edit GET {path.replace(B, '')}: vuelve al listado con '{msg}'"), st == 200 and any(msg in t for _, t in parse_messages(page)[1]), f"status={st} {parse_messages(page)[1][:1]}")

    def f_(lab, art, qty, rid=0, rv="", ret="", flt="", term="", upd=None, drop=()):
        f = [("Input.Id", rid), ("Input.LaboratoryId", lab), ("Input.ArticleId", art), ("Input.RowVersionBase64", rv), ("Input.LaboratoryFilter", flt), ("Input.SearchTerm", term), ("ReturnUrl", ret)]
        if qty is not None:
            f.append(("Input.AvailableQuantity", qty))
        if upd is not None:
            f.append(("Input.UpdatedOn", upd))
        return [x for x in f if x[0] not in drop]

    def epost(lab, art, fields, name, follow=True):
        return c.raw_post(f"{B}/Edit/{lab}/{art}", fields, follow=follow, token_from=f"{B}/Edit/{lab}/{art}", name=name)

    # ---------------------------------------------------------------- Edit: valores límite sobre la fila nueva
    max_msgs = {}
    # cada valor inválido se prueba sobre su propio par (laboratorio, artículo) para que un valor aceptado por error no contamine los demás
    probe_labs = labs[3:3 + 7]
    for (q, desc, msg), lab_i in zip((("-1", "-1", "La cantidad debe ser mayor o igual a 0."), ("-0.01", "-0.01", "La cantidad debe ser mayor o igual a 0."),
                                      ("abc", "'abc'", "El valor 'abc' no es un número válido."), ("1e3", "'1e3'", "no es un número válido."),
                                      ("999999.01", "999999.01", "La cantidad debe"), ("1000000", "1000000", "La cantidad debe"),
                                      ("999999.999", "999999.999", "La cantidad debe")), probe_labs):
        r = epost(lab_i, A0, f_(lab_i, A0, q), f"qty {desc}")
        max_msgs[q] = r.errors
        expect(R, L(f"Edit: cantidad {desc}"), r, status=200, error=msg)
        created = stock(lab_i, A0)
        check(R, L(f"Edit: cantidad {desc} no crea la fila"), created is None, f"fila creada con {created[1] if created else None}")
    check(R, L("Edit: el mensaje para una cantidad por encima del máximo (1000000) menciona el límite 999999"), any("999" in e for e in max_msgs["1000000"]),
          f"mensaje mostrado={max_msgs['1000000'][:1]} (Range(0,999999) reutiliza «mayor o igual a 0»)")
    # vacío / ausente
    LAB3 = labs[2]
    r = epost(LAB2, A0, f_(LAB2, A0, ""), "qty vacía")
    check(R, L("Edit: cantidad vacía se rechaza en vez de guardarse en silencio como 0 [sospecha del plan]"), r.status == 200 and stock(LAB2, A0) is None,
          f"status={r.status} errores={r.errors[:2]} fila={stock(LAB2, A0)}")
    r = epost(LAB3, A0, f_(LAB3, A0, None), "qty ausente")
    check(R, L("Edit: cantidad ausente del POST se rechaza en vez de guardarse en silencio como 0"), r.status == 200 and stock(LAB3, A0) is None,
          f"status={r.status} errores={r.errors[:2]} fila={stock(LAB3, A0)}")
    # fecha de actualización
    r = epost(LAB1, A0, f_(LAB1, A0, "5", upd="no-es-fecha"), "fecha inválida")
    check(R, L("Edit: fecha de actualización inválida se rechaza sin 500 (y sin crear la fila)"), r.status == 200 and stock(LAB1, A0) is None and bool(r.errors), f"status={r.status} errores={r.errors[:2]}")

    # ---------------------------------------------------------------- Edit: crear
    r = epost(LAB1, A0, f_(LAB1, A0, "12,5", upd="2026-10-16", term=A0N, flt=LAB1), "crear 12,5", follow=False)
    check(R, L("Edit: crear con '12,5' redirige al listado con los filtros"), r.status == 302 and (r.location or "").startswith("/Catalogs/Stock") and f"LaboratoryFilter={LAB1}".lower() in (r.location or "").lower(),
          f"status={r.status} location={r.location}")
    s = stock(LAB1, A0)
    check(R, L("Edit: fila creada con 12.50 y fecha"), s and float(s[1]) == 12.5 and s[3] is not None, str(s))
    r = epost(LAB1, A0, f_(LAB1, A0, "3"), "crear de nuevo")
    expect(R, L("Edit: crear de nuevo la misma fila (Id=0)"), r, status=200, error="Ya existe una existencia para este artículo en este laboratorio. Recargue la página para editarla.")
    check(R, L("Edit: la fila duplicada no se creó ni cambió"), float(stock(LAB1, A0)[1]) == 12.5)

    # mensaje de éxito
    r = epost(LAB1, A0, f_(LAB1, A0, "14", rid=s[0], rv=base64.b64encode(s[2]).decode()), "actualizar 14")
    expect(R, L("Edit: actualizar la fila"), r, status=302, toast=f"Existencia actualizada para «{A0N}» en «{lab1_code} - {lab1_name}».", toast_kind="success")
    s = stock(LAB1, A0)
    check(R, L("Edit: BD 14.00 y RowVersion cambió"), float(s[1]) == 14, str(s))

    # límites aceptados
    for q, exp, desc in (("0", 0.0, "0"), ("999999", 999999.0, "999999 (máximo)"), ("0.01", 0.01, "0.01"), ("12.34", 12.34, "12.34"), ("1,5", 1.5, "'1,5'"), ("  7  ", 7.0, "'  7  '")):
        s = stock(LAB1, A0)
        r = epost(LAB1, A0, f_(LAB1, A0, q, rid=s[0], rv=base64.b64encode(s[2]).decode()), f"límite {desc}")
        got = float(stock(LAB1, A0)[1])
        check(R, L(f"Edit: cantidad {desc} se acepta y queda en {exp}"), r.status == 302 and abs(got - exp) < 0.0001, f"status={r.status} errores={r.errors[:2]} BD={got}")
    # más de 2 decimales: la columna es decimal(12,2); el Kit y el Import limitan a 2 decimales, aquí se debería rechazar
    for q in ("0.001", "1.234", "0.005"):
        s = stock(LAB1, A0)
        before = float(s[1])
        r = epost(LAB1, A0, f_(LAB1, A0, q, rid=s[0], rv=base64.b64encode(s[2]).decode()), f"decimales {q}")
        got = float(stock(LAB1, A0)[1])
        check(R, L(f"Edit: cantidad {q} (3 decimales) se rechaza en vez de redondearse en silencio"), r.status == 200 and bool(r.errors),
              f"status={r.status} errores={r.errors[:1]} BD {before} -> {got} (Kit e Import sí limitan a 2 decimales)")
    # ---------------------------------------------------------------- Edit: concurrencia
    s = stock(LAB1, A0)
    rv_old = base64.b64encode(s[2]).decode()
    r = epost(LAB1, A0, f_(LAB1, A0, "20", rid=s[0], rv=rv_old), "otra persona edita")
    check(R, L("Edit: 1.er guardado con RowVersion vigente funciona"), r.status == 302 and float(stock(LAB1, A0)[1]) == 20)
    r = epost(LAB1, A0, f_(LAB1, A0, "99", rid=s[0], rv=rv_old), "RowVersion vieja", follow=False)
    check(R, L("Edit: guardar con RowVersion vieja -> redirige a Edit (recarga)"), r.status == 302 and f"/Catalogs/Stock/Edit/{LAB1}/{A0}" in (r.location or ""), f"status={r.status} loc={r.location}")
    st, page = c.get((r.location or "/"))
    check(R, L("Edit: aviso de concurrencia 'Otra persona modificó esta existencia; se recargaron los valores actuales.'"),
          any("Otra persona modificó esta existencia; se recargaron los valores actuales." in t and k == "warning" for k, t in parse_messages(page)[1]), str(parse_messages(page)[1][:2]))
    check(R, L("Edit: la cantidad no se sobrescribió (sigue 20)"), float(stock(LAB1, A0)[1]) == 20)
    r = epost(LAB1, A0, f_(LAB1, A0, "30", rid=s[0], rv=""), "sin RowVersion")
    check(R, L("Edit: sin RowVersion en el POST el guardado se acepta SIN control de concurrencia [hallazgo extra]"), float(stock(LAB1, A0)[1]) == 20,
          f"status={r.status}; cantidad en BD={float(stock(LAB1, A0)[1])} (el servidor omite la comprobación si falta el token)")
    s = stock(LAB1, A0)
    r = epost(LAB1, A0, f_(LAB1, A0, "31", rid=s[0], rv="@@no-es-base64@@"), "RowVersion malformada")
    check(R, L("Edit: RowVersion malformada tampoco provoca 500 ni sobrescribe sin control [hallazgo extra]"), r.status in (200, 302) and float(stock(LAB1, A0)[1]) != 31,
          f"status={r.status}; cantidad en BD={float(stock(LAB1, A0)[1])}")
    # fila ajena: Id de otra fila (QAC) con LaboratoryId/ArticleId de la URL
    s = stock(LAB1, A0)
    epost(LAB2, arts[4][0], f_(LAB2, arts[4][0], "5"), "fila ajena (preparación)")
    other = stock(LAB2, arts[4][0])[0]
    before_other = float(stock(LAB2, arts[4][0])[1])
    r = epost(LAB1, A0, f_(LAB1, A0, "123", rid=other, rv=""), "Id ajeno")
    after_other = float(stock(LAB2, arts[4][0])[1])
    check(R, L("Edit: un Input.Id de OTRA fila (otro laboratorio/artículo) no modifica esa fila [hallazgo extra: el POST usa el Id del cuerpo, no la ruta]"),
          after_other == before_other, f"fila {other}: {before_other} -> {after_other}; toasts={r.toasts[:1]}")
    r = epost(LAB1, A0, f_(LAB1, A0, "5", rid=99999999), "Id inexistente")
    expect(R, L("Edit: Input.Id de una fila que ya no existe"), r, status=302, toast="La existencia que intenta modificar ya no existe.", toast_kind="error")

    # ---------------------------------------------------------------- Edit: laboratorio / artículo inválidos en el POST
    r = epost(LAB1, A0, f_(99999, A0, "1"), "lab inexistente")
    expect(R, L("Edit POST: laboratorio inexistente"), r, status=200, error="El laboratorio seleccionado no existe o no está activo.")
    r = epost(LAB1, A0, f_(LAB1, 99999999, "1"), "art inexistente")
    expect(R, L("Edit POST: artículo inexistente"), r, status=200, error="El artículo seleccionado no existe o no está activo.")
    c.submit(f"/Catalogs/Articles/Edit/{arts[5][0]}", {"Input.Status": "1"})
    r = epost(LAB1, arts[5][0], f_(LAB1, arts[5][0], "1"), "art inactivo")
    expect(R, L("Edit POST: artículo inactivo"), r, status=200, error="El artículo seleccionado no existe o no está activo.")
    st, page = c.get(f"{B}/Edit/{LAB1}/{arts[5][0]}")
    check(R, L("Edit GET: artículo inactivo vuelve al listado con el error"), any("El artículo seleccionado no existe o no está activo." in t for _, t in parse_messages(page)[1]))
    val_before = float(stock(LAB1, A0)[1])
    tokenless(R, c, L("Edit"), f"{B}/Edit/{LAB1}/{A0}", f_(LAB1, A0, "77", rid=stock(LAB1, A0)[0]))
    check(R, L("POST sin token no modificó nada"), float(stock(LAB1, A0)[1]) == val_before)

    # ---------------------------------------------------------------- Edit: returnUrl
    s = stock(LAB1, A0)
    rv = lambda: base64.b64encode(stock(LAB1, A0)[2]).decode()  # noqa: E731
    for ret, expect_loc, desc in (("/Catalogs/Stock?LaboratoryFilter=2", "/Catalogs/Stock?LaboratoryFilter=2", "ruta local con query"), ("/LaboratoryOperations/Requerimientos?planId=1", "/LaboratoryOperations/Requerimientos?planId=1", "ruta local de otra pantalla")):
        r = epost(LAB1, A0, f_(LAB1, A0, "21", rid=s[0], rv=rv(), ret=ret), f"ret {desc}", follow=False)
        check(R, L(f"Edit: returnUrl local ({desc}) se respeta"), r.status == 302 and r.location == expect_loc, f"location={r.location}")
    for ret in ("https://evil.example/x", "//evil.example/x", "http://evil.example", "/\\evil.example", "javascript:alert(1)", "evil.example", "\\\\evil.example"):
        r = epost(LAB1, A0, f_(LAB1, A0, "22", rid=s[0], rv=rv(), ret=ret), f"ret {ret}", follow=False)
        loc = r.location or ""
        check(R, L(f"Edit: returnUrl externo/inválido {ret!r} se ignora"), r.status == 302 and loc.startswith("/Catalogs/Stock") and "evil" not in loc, f"location={loc}")
    st, page = c.get(f"{B}/Edit/{LAB1}/{A0}?returnUrl={urllib.parse.quote('https://evil.example/x')}")
    check(R, L("Edit GET con returnUrl externo no lo refleja en los enlaces"), st == 200 and "evil.example" not in page, f"status={st}")
    evil = urllib.parse.quote('"><script>alert(1)</script>')
    st, page = c.get(f"{B}/Edit/{LAB1}/{A0}?returnUrl={evil}")
    check(R, L("Edit GET con returnUrl con HTML no lo refleja sin escapar"), st == 200 and "<script>alert(1)</script>" not in page, f"status={st}")

    # ---------------------------------------------------------------- Import: GET, plantilla
    st, page = c.get(f"{B}/Import")
    check(R, L("Import GET responde 200 con el formulario de carga"), st == 200 and 'name="Upload"' in page, f"status={st}")
    op = c.opener
    resp = op.open(f"http://localhost:5000{B}/Import?handler=Template", timeout=300)
    data = resp.read()
    cd = resp.headers.get("Content-Disposition", "")
    check(R, L("Import/Template: descarga 'Existencias_yyyyMMdd.xlsx' con tipo Excel"),
          resp.status == 200 and XLSX in resp.headers.get("Content-Type", "") and re.search(r"Existencias_\d{8}\.xlsx", cd) is not None, f"ct={resp.headers.get('Content-Type')} cd={cd}")
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data))
        ws = wb["ExistenciasArticulos"]
        hdr = [ws.cell(1, i).value for i in range(1, 4)]
        nrows = ws.max_row - 1
        check(R, L("Import/Template: hoja 'ExistenciasArticulos' con LaboratorioCodigo, ArticuloCodigo, CantidadDisponible"), hdr == HEAD, str(hdr))
        check(R, L("Import/Template: una fila por cada existencia de la base"), nrows == scalar("select count(*) from ArticleStocks s join Laboratories l on l.Id=s.LaboratoryId and l.Status<>2 join Articles a on a.Id=s.ArticleId and a.Status<>2"), f"filas={nrows}")
    except Exception as exc:
        check(R, L("Import/Template: el archivo es un Excel legible"), False, f"{type(exc).__name__}: {exc}")

    # ---------------------------------------------------------------- Import: vista previa, archivos inválidos
    IM = f"{B}/Import"
    PV = f"{IM}?handler=Preview"

    def pv(content, filename="datos.xlsx", name="preview", ctype=XLSX, field="Upload"):
        files = [] if content is None else [(field, filename, content, ctype)]
        return multipart_post(c, PV, [], files, token_from=IM, name=name)

    r = pv(None, name="sin archivo")
    expect(R, L("Import/Preview: sin archivo"), r, status=302, toast="Selecciona un archivo Excel (.xlsx).", toast_kind="error")
    r = pv(b"", name="archivo vacío")
    expect(R, L("Import/Preview: archivo vacío (0 bytes)"), r, status=302, toast="Selecciona un archivo Excel (.xlsx).", toast_kind="error")
    r = pv(b"a,b,c\n1,2,3\n", filename="datos.csv", ctype="text/csv", name="csv")
    expect(R, L("Import/Preview: archivo .csv"), r, status=302, toast="El archivo debe ser .xlsx y pesar menos de 10 MB.", toast_kind="error")
    r = pv(xlsx_bytes([["H-01", "X", 1]]), filename="datos.xls", name="xls")
    expect(R, L("Import/Preview: extensión .xls"), r, status=302, toast="El archivo debe ser .xlsx y pesar menos de 10 MB.", toast_kind="error")
    r = pv(b"a,b,c\n1,2,3\n", filename="renombrado.xlsx", name="csv renombrado")
    expect(R, L("Import/Preview: .csv renombrado a .xlsx"), r, status=302, toast="No se pudo leer el archivo. Verifica que sea un Excel válido.", toast_kind="error")
    r = pv(b"\x00" * 200, filename="ceros.xlsx", name="basura")
    expect(R, L("Import/Preview: bytes basura con extensión .xlsx"), r, status=302, toast="No se pudo leer el archivo. Verifica que sea un Excel válido.", toast_kind="error")
    r = pv(b"PK\x03\x04" + b"\x00" * (10 * 1024 * 1024 + 10), filename="grande.xlsx", name="más de 10 MB")
    expect(R, L("Import/Preview: archivo de más de 10 MB"), r, status=302, toast="El archivo debe ser .xlsx y pesar menos de 10 MB.", toast_kind="error")
    r = pv(xlsx_bytes([["H-01", "X", 1]], header=["LaboratorioCodigo", "ArticuloCodigo"]), name="falta una columna")
    expect(R, L("Import/Preview: falta la columna CantidadDisponible"), r, status=302, toast="Faltan columnas: CantidadDisponible. Descarga la plantilla para ver el formato.", toast_kind="error")
    r = pv(xlsx_bytes([["a", "b", "c"]], header=["Lab", "Art", "Cant"]), name="faltan todas")
    expect(R, L("Import/Preview: faltan las tres columnas"), r, status=302, toast="Faltan columnas: LaboratorioCodigo, ArticuloCodigo, CantidadDisponible.", toast_kind="error")
    r = pv(xlsx_bytes([], header=None), name="hoja vacía")
    expect(R, L("Import/Preview: hoja completamente vacía"), r, status=302, toast="La hoja está vacía.", toast_kind="error")
    r = pv(xlsx_bytes([]), name="solo encabezados")
    expect(R, L("Import/Preview: solo encabezados, sin filas"), r, status=302, toast="El archivo no tiene filas con datos.", toast_kind="error")
    r = pv(xlsx_bytes([[None, None, None], ["", "", ""]]), name="filas en blanco")
    expect(R, L("Import/Preview: solo filas en blanco"), r, status=302, toast="El archivo no tiene filas con datos.", toast_kind="error")
    r = pv(xlsx_bytes([[lab1_code, imp[0][2], 1]] * 1 + [["H-01", f"ART-X{i}", 1] for i in range(5000)]), name="5001 filas")
    expect(R, L("Import/Preview: 5001 filas"), r, status=302, toast="El archivo supera el máximo de 5000 filas.", toast_kind="error")
    r = pv(xlsx_bytes([[lab1_code, imp[0][2], 1]] + [[lab1_code, "ART-NOEXISTE", 1]] * 4999), name="5000 filas exactas")
    n, u, un, e, tok, rows = preview_summary(r.text) if r.status == 200 else (None,) * 6
    check(R, L("Import/Preview: 5000 filas exactas se aceptan (1 nueva, 4999 con error por repetidas/inexistentes)"), r.status == 200 and n == 1 and e == 4999, f"status={r.status} nuevas={n} errores={e}")

    # ---------------------------------------------------------------- Import: filas con error y vista previa
    c1, c2, c3, c4 = imp[0][2], imp[1][2], imp[2][2], imp[3][2]
    rows_x = [
        [lab1_code, c1, 10],                      # nueva
        [lab1_code.lower(), c2.lower(), "7,5"],   # nueva (minúsculas y coma)
        ["ZZZ-NO", c3, 3],                        # laboratorio inexistente
        [lab1_code, "ART-NOEXISTE", 3],           # artículo inexistente
        [lab1_code, c3, "abc"],                   # cantidad no numérica
        [lab1_code, c3, -1],                      # negativa
        [lab1_code, c3, 1000000],                 # mayor al máximo
        [lab1_code, c3, None],                    # cantidad vacía
        [lab1_code, c1, 99],                      # repetida
        [None, c3, 4],                            # sin laboratorio
        [lab1_code, None, 4],                     # sin artículo
        [lab1_code, c4, 0],                       # nueva con 0
        [lab1_code, A0C, float(stock(LAB1, A0)[1])],   # igual a la existente (sin cambios)
        [lab1_code, A0C.lower(), 21],             # repetida (misma fila en minúsculas)
    ]
    r = pv(xlsx_bytes(rows_x), name="mezcla")
    n, u, un, e, tok, rows = preview_summary(r.text) if r.status == 200 else (None,) * 6
    check(R, L("Import/Preview: mezcla -> 3 nuevas, 0 a actualizar, 1 sin cambios, 10 con error"), r.status == 200 and (n, u, un, e) == (3, 0, 1, 10), f"status={r.status} (n,u,un,e)={(n, u, un, e)}")
    txt = " | ".join(t for _, t in rows)
    for frag in ("Laboratorio ZZZ-NO inexistente o inactivo.", "Artículo ART-NOEXISTE inexistente o inactivo.", "Cantidad vacía o no numérica.",
                 "La cantidad debe estar entre 0 y 999.999.", "Fila repetida (mismo laboratorio y artículo).", "Falta el laboratorio o el artículo."):
        check(R, L(f"Import/Preview: mensaje por fila «{frag}»"), frag in txt, "no aparece")
    check(R, L("Import/Preview: la vista previa no escribe nada en la base"), stock(LAB1, imp[0][0]) is None and tok is not None)
    first_tok = tok

    # ---------------------------------------------------------------- Import: aplicar
    def apply(token, name="apply"):
        return c.raw_post(f"{IM}?handler=Apply", [("Token", token)] if token is not None else [], token_from=IM, name=name)

    r = apply(None, "sin token")
    expect(R, L("Import/Apply: sin token de vista previa"), r, status=302, toast="La vista previa expiró. Vuelve a subir el archivo.", toast_kind="error")
    r = apply("no-existe", "token falso")
    expect(R, L("Import/Apply: token inexistente"), r, status=302, toast="La vista previa expiró. Vuelve a subir el archivo.", toast_kind="error")
    tokenless(R, c, L("Import/Apply"), f"{IM}?handler=Apply", [("Token", first_tok)])
    r = apply(first_tok)
    expect(R, L("Import/Apply: aplicar la mezcla (solo se escriben las 3 nuevas)"), r, status=302, redirect=B, toast="Existencias cargadas: 3 nueva(s) y 0 actualizada(s).", toast_kind="success")
    check(R, L("Import/Apply: filas escritas con 10, 7.5 y 0; la 'sin cambios' no se tocó"),
          stock(LAB1, imp[0][0]) and float(stock(LAB1, imp[0][0])[1]) == 10 and float(stock(LAB1, imp[1][0])[1]) == 7.5 and float(stock(LAB1, imp[3][0])[1]) == 0 and stock(LAB1, imp[2][0]) is None,
          f"{[stock(LAB1, a[0]) for a in imp[:4]]}")
    r = apply(first_tok, "doble clic")
    expect(R, L("Import/Apply: segundo clic con el mismo token"), r, status=302, toast="La vista previa expiró. Vuelve a subir el archivo.", toast_kind="error")
    check(R, L("Import/Apply: el segundo clic no duplicó filas"), scalar("select count(*) from ArticleStocks where LaboratoryId=? and ArticleId in (?,?,?)", LAB1, imp[0][0], imp[1][0], imp[3][0]) == 3)

    # volver a subir: ahora hay actualizaciones y sin cambios
    r = pv(xlsx_bytes([[lab1_code, c1, 10], [lab1_code, c2, 8], [lab1_code, c4, 5], [lab1_code, c3, 2]]), name="segunda carga")
    n, u, un, e, tok, rows = preview_summary(r.text)
    check(R, L("Import/Preview: segunda carga -> 1 nueva, 2 a actualizar, 1 sin cambios"), (n, u, un, e) == (1, 2, 1, 0), f"{(n, u, un, e)}")
    r = apply(tok)
    expect(R, L("Import/Apply: segunda carga (las 'sin cambios' no se reescriben)"), r, status=302, redirect=B, toast="Existencias cargadas: 1 nueva(s) y 2 actualizada(s).", toast_kind="success")
    # solo sin cambios / solo errores: no hay nada que aplicar
    r = pv(xlsx_bytes([[lab1_code, c1, 10]]), name="solo sin cambios")
    n, u, un, e, tok, rows = preview_summary(r.text)
    check(R, L("Import/Preview: solo 'sin cambios' -> sin botón de aplicar"), (n, u, un, e) == (0, 0, 1, 0) and tok is None, f"{(n, u, un, e)} token={tok}")
    r = pv(xlsx_bytes([["ZZZ", c1, 10]]), name="solo errores")
    n, u, un, e, tok, rows = preview_summary(r.text)
    check(R, L("Import/Preview: solo errores -> sin botón de aplicar"), (n, u, un, e) == (0, 0, 0, 1) and tok is None, f"{(n, u, un, e)} token={tok}")
    # token de una vista previa sin cambios aplicado a mano
    # (el servidor guarda la sesión aunque no haya cambios: se aplica 'No había cambios')
    # encabezados en otro orden, otra hoja y mayúsculas/minúsculas distintas
    wb = openpyxl.Workbook()
    ws0 = wb.active
    ws0.title = "Otra"
    ws0.append(["x"])
    ws = wb.create_sheet("ExistenciasArticulos")
    ws.append(["cantidaddisponible", "ARTICULOCODIGO", "laboratoriocodigo", "extra"])
    ws.append([6, c2, lab1_code, "ignorada"])
    bio = io.BytesIO()
    wb.save(bio)
    r = pv(bio.getvalue(), name="orden de columnas y hoja")
    n, u, un, e, tok, rows = preview_summary(r.text)
    check(R, L("Import/Preview: usa la hoja 'ExistenciasArticulos', columnas en otro orden/caso y con columnas extra"), (n, u, un, e) == (0, 1, 0, 0), f"{(n, u, un, e)}")
    # lotes de 100: 5 artículos QAC x todos los laboratorios activos
    big = [[scalar("select Code from Laboratories where Id=?", lb), a[2], 3] for lb in labs for a in arts[:5]]
    r = pv(xlsx_bytes(big), name="lotes")
    n, u, un, e, tok, rows = preview_summary(r.text)
    check(R, L(f"Import/Preview: {len(big)} filas (más de 100) válidas"), e == 0 and (n + u + un) == len(big) and len(big) > 100, f"{(n, u, un, e)} total={len(big)}")
    r = apply(tok)
    toast = r.toasts[0][1] if r.toasts else ""
    m = re.search(r"(\d+) nueva\(s\) y (\d+) actualizada\(s\)", toast)
    created, updated = (int(m.group(1)), int(m.group(2))) if m else (None, None)
    total_db = scalar("select count(*) from ArticleStocks where ArticleId in (?,?,?,?,?) and AvailableQuantity=3", *[a[0] for a in arts[:5]])
    check(R, L("Import/Apply: lotes de 100 en una transacción -> todas las filas del archivo quedan con 3"), m is not None and r.status == 302 and total_db == len(big),
          f"toast={toast!r} filas con 3 en BD={total_db} de {len(big)}")
    check(R, L("Import/Apply: nuevas + actualizadas = filas que cambiaban"), m is not None and created + updated == n + u, f"{toast!r} esperado {n}+{u}")
    tokenless(R, c, L("Import/Preview"), PV, [])
    check(R, L("Total de existencias de la base solo creció por las filas QAC"),
          scalar("select count(*) from ArticleStocks") - stock0 == scalar("select count(*) from ArticleStocks where ArticleId in (?,?,?,?,?,?)", *[a[0] for a in arts]),
          f"antes={stock0} ahora={scalar('select count(*) from ArticleStocks')}")

    # ---------------------------------------------------------------- limpieza: artículos QAC dados de baja (las existencias no se borran desde la UI)
    for a in arts:
        c.submit(f"/Catalogs/Articles/Delete/{a[0]}", {})
    left = scalar("select count(*) from Articles where Name like 'QAC Stk%' and Status<>2")
    check(R, L("Limpieza: artículos QAC de existencias dados de baja"), left == 0, f"quedan={left}")


if __name__ == "__main__":
    run_module("stock", run)
