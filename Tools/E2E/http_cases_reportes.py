"""Reportes Excel por HTTP: descargas válidas (xlsx abre con openpyxl y no está vacío) y errores esperados. Solo contra la copia QA."""
import io
import os
import sys

import openpyxl

import e2e_lib
from http_forms import Client

PW = os.environ["E2E_PW"]
c = Client(os.environ.get("E2E_USER", "qa_super"), PW)
R = []


def fetch(path):
    import urllib.error
    import urllib.request
    try:
        r = c.opener.open(e2e_lib.BASE + path, timeout=180)
        return r.status, r.headers.get("Content-Type", ""), r.headers.get("Content-Disposition", ""), r.read()
    except urllib.error.HTTPError as e:
        return e.code, e.headers.get("Content-Type", ""), "", e.read()


def ok_xlsx(name, path):
    st, ct, cd, body = fetch(path)
    detail = f"status={st} ct={ct} bytes={len(body)} cd={cd[:70]}"
    good = False
    if st == 200 and len(body) > 1000:
        try:
            wb = openpyxl.load_workbook(io.BytesIO(body))
            cells = sum(1 for ws in wb for row in ws.iter_rows() for cell in row if cell.value not in (None, ""))
            detail += f" hojas={len(wb.sheetnames)} celdas_con_datos={cells}"
            good = cells > 5
        except Exception as ex:  # archivo corrupto
            detail += f" ERROR_XLSX={type(ex).__name__}"
    R.append(e2e_lib.Result(name, "xlsx válido con datos", "xlsx válido con datos" if good else "NO válido", detail))


def err(name, path, status, fragment=None):
    st, ct, cd, body = fetch(path)
    text = body.decode("utf-8", "replace")
    good = st == status and (fragment is None or fragment.lower() in text.lower())
    expected = f"{status}" + (f" '{fragment}'" if fragment else "")
    R.append(e2e_lib.Result(name, expected, expected if good else f"{st} '{text[:120]}'", ""))


MGMT, LAB, UNIT, REQ, PLAN, DEP, ACQ = 30, 3, 382, 330, 2, 2, 331
# ids de datos reales (por SQL de solo lectura); si no existen, la comprobación se omite
UNIT_L6 = e2e_lib.db_scalar("SELECT TOP 1 v.EquipmentUnitId FROM Verifications v "
                            "JOIN EquipmentUnits u ON u.Id = v.EquipmentUnitId "
                            "WHERE v.ManagementId = ? AND u.LaboratoryId = ? ORDER BY v.Id", MGMT, LAB)
REQ_COMPRA = e2e_lib.db_scalar("SELECT TOP 1 Id FROM Requests WHERE Type = ? ORDER BY Id", 2)  # RequestType.Purchasing = 2


def omitida(motivo):
    R.append(e2e_lib.Result(f"omitida: sin datos ({motivo})", "omitida", "omitida", ""))


ok_xlsx("L-6 consolidado (laboratorio)", f"/api/reports/download/l6?managementId={MGMT}&labId={LAB}")
if UNIT_L6:
    ok_xlsx("L-6 por unidad", f"/api/reports/download/l6?managementId={MGMT}&labId={LAB}&unitId={UNIT_L6}")
else:
    omitida("L-6 por unidad: sin Verifications en gestion/laboratorio")
ok_xlsx("Plan L-48", f"/api/reports/download/l48?managementId={MGMT}&labId={LAB}")
ok_xlsx("L-7 solicitud", f"/api/reports/download/l7?requestId={REQ}&managementId={MGMT}")
ok_xlsx("L-8 Kardex", f"/api/reports/download/l8?managementPlanId={PLAN}")
ok_xlsx("L-3 salida", f"/api/reports/download/l3?departureId={DEP}")
if REQ_COMPRA:
    ok_xlsx("L-12 adquisición", f"/api/reports/download/adquisicion?requestId={REQ_COMPRA}")
else:
    omitida("L-12 adquisición: sin Requests de compra (Type=Purchasing)")

err("L-6 gestión inexistente", f"/api/reports/download/l6?managementId=99999&labId={LAB}", 404, "Gestión no encontrada")
err("L-6 sin datos para el equipo", f"/api/reports/download/l6?managementId={MGMT}&labId={LAB}&unitId=1", 404, "No hay datos L-6")
err("L-8 sin parámetros", "/api/reports/download/l8", 400, "managementPlanId o unitId")
err("L-7 sin parámetros", "/api/reports/download/l7", 400)
err("Adquisición sin parámetros", "/api/reports/download/adquisicion", 400)
err("L-3 equipo sin salidas", f"/api/reports/download/l3?unitId=1&managementId={MGMT}", 404)

sys.exit(1 if e2e_lib.run_all([lambda: R]) else 0)
