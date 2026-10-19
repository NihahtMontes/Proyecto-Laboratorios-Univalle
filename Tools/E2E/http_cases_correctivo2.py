"""
Flujo de gestion CORRECTIVA semestral - parte 2 (CR-08 a CR-15, sin CR-14). Reutiliza los helpers de
http_cases_correctivo_common (C) y http_cases_procesos_common (P). Solo contra la copia QA.

- CR-08 «Rehacer fase» (POST /api/WizardRollback/redo) sobre planes correctivos: L-6 rechazado, L-3 con
  salida cancelada y completado bloqueado por D-7 (otro proceso abierto).
- CR-09 contadores del dashboard de la gestion correctiva vs SQL.
- CR-10 informes Excel en correctivo (L-6 rechazado; L-7/L-8/L-3/adquisicion/L-48 validos).
- CR-11 permisos: Supervisor entra por el asistente correctivo pero no a Maintenances/Create;
  Cierre de practicas no accede a /Requests/Create.
- CR-12 concurrencia: dos L-7 simultaneos para la misma unidad -> max 1 Request y 1 plan (D-7 + indice unico).
- CR-13 gestion correctiva NO activa: L-7 rechazado con warning y sin cambios.
- CR-15 convivencia: cronograma preventivo/correctivo y edicion de solicitud correctiva.

Salida: outputs/qa-usuario/http_cases_procesos_correctivo2.txt (la escribe P.finish).
"""
import io
import os
import re
import threading
import time

import openpyxl

import http_cases_procesos_common as P
import http_cases_correctivo_common as C
from http_cases_procesos_common import check, plan, scalar, sql

REDO_URL = "/api/WizardRollback/redo"       # Controllers/WizardRollbackController.cs [Route("api/[controller]")] + RedoPhase
TOKEN_PAGE = "/Equipment/Create"            # pagina cualquiera del sitio para el token antifalsificacion
LOAN_CANCELLED = 99                         # LoanStatus.Cancelled
PHASE_L6, PHASE_L7, PHASE_L8, PHASE_L3, PHASE_KARDEX, PHASE_L12 = 1, 2, 3, 4, 5, 6   # WizardPhase


def redo(cl, plan_id, target):
    """POST /api/WizardRollback/redo {planId, targetPhase}; devuelve (status, cuerpo)."""
    return P.json_post(cl, REDO_URL, {"planId": plan_id, "targetPhase": target}, token_page=TOKEN_PAGE)


# ---------------------------------------------------------------- reportes (imita http_cases_reportes.fetch)
def fetch(cl, path):
    import urllib.error
    import urllib.request
    op = urllib.request.build_opener(urllib.request.HTTPCookieProcessor(cl.jar))
    try:
        r = op.open(P.BASE + path, timeout=180)
        return r.status, r.read()
    except urllib.error.HTTPError as e:
        return e.code, e.read()


def ok_xlsx(R, nombre, cl, path):
    st, body = fetch(cl, path)
    det = f"status={st} bytes={len(body)}"
    good = False
    if st == 200 and len(body) > 500:
        try:
            wb = openpyxl.load_workbook(io.BytesIO(body))
            cells = sum(1 for ws in wb for row in ws.iter_rows() for cell in row if cell.value not in (None, ""))
            det += f" hojas={len(wb.sheetnames)} celdas_con_datos={cells}"
            good = cells > 3
        except Exception as ex:
            det += f" ERROR_XLSX={type(ex).__name__}"
    check(R, nombre, good, det)


def err_report(R, nombre, cl, path, status, fragmento):
    st, body = fetch(cl, path)
    text = body.decode("utf-8", "replace")
    check(R, nombre, st == status and fragmento.lower() in text.lower(), f"status={st} '{text[:120]}'")


def run(R, c):
    pw = os.environ["E2E_PW"]
    lab = P.lab_propio(c, "correctivo2")
    corr = C.asegurar_correctiva(c)

    # ================================================================ CR-08 «Rehacer fase» en correctivo
    ctxa = C.llevar_corr(c, corr, lab, "c8a", hasta="L7")        # plan 3/4 (L-7 registrado)
    ctxb = C.llevar_corr(c, corr, lab, "c8b", hasta="KARDEX")    # plan 6/8: L-7/L-8/L-3/Kardex hechos
    ctxc = C.llevar_corr(c, corr, lab, "c8c", hasta="DONE")      # plan 6/9 PlanStatus 2 (Completado)
    dep_b = plan(ctxb["plan"])["dep"]
    r2c = C.l7_corr(c, corr, lab, ctxc["unit"], name="CR-08 segundo proceso abierto de la unidad", follow=True)
    pid_c2 = C.plan_de(corr, ctxc["unit"])
    check(R, "CR-08 fixture: la unidad de c8c queda con 2 planes (uno Completado y otro abierto 3/4)",
          r2c.status == 302 and pid_c2 and pid_c2 != ctxc["plan"]
          and plan(ctxc["plan"])["pstatus"] == 2 and plan(pid_c2)["phase"] == 3,
          f"r={P.dump(r2c)} p1={plan(ctxc['plan'])} p2={plan(pid_c2)}")

    # (a) rehacer a L-6 (TargetPhase 1) en correctivo -> 400 y sin cambios
    st, body = redo(c, ctxa["plan"], PHASE_L6)
    check(R, "CR-08 (a) rehacer a L-6 (TargetPhase=1) en correctivo -> 400 'no forma parte de una gestion correctiva'",
          st == 400 and "no forma parte de una gesti\u00f3n correctiva" in body.replace('"', ""),
          f"status={st} {body[:120]}")
    check(R, "CR-08 (a) el plan queda intacto (sigue 3/4 esperando L-8)",
          plan(ctxa["plan"])["phase"] == 3 and plan(ctxa["plan"])["state"] == 4, f"plan={plan(ctxa['plan'])}")

    # (b) plan en Kardex rehecho a L-3 (TargetPhase 4): vuelve a fase 4 y la salida queda cancelada.
    # El fin del L-8 es "el minuto siguiente" y el rollback cierra historiales en UtcNow: esperar como el modulo de rollback.
    time.sleep(75)
    st, body = redo(c, ctxb["plan"], PHASE_L3)
    p = plan(ctxb["plan"])
    check(R, "CR-08 (b) rehacer a L-3 (TargetPhase=4) un plan en Kardex -> 200 'Progreso revertido'",
          st == 200 and "Progreso revertido" in body, f"status={st} {body[:120]}")
    check(R, "CR-08 (b) el plan vuelve a la fase 4 (Exit/AwaitingDeparture 4/6) con la salida desvinculada y Cancelada (99)",
          p["phase"] == 4 and p["state"] == 6 and p["dep"] is None
          and scalar("select Status from Departures where Id=?", dep_b) == LOAN_CANCELLED,
          f"plan={p} salida={dep_b}->{scalar('select Status from Departures where Id=?', dep_b)}")

    # (c) completado con otro proceso abierto en la misma unidad -> 400 D-7 y nada cambia
    antes = (plan(ctxc["plan"]), plan(pid_c2))
    st, body = redo(c, ctxc["plan"], PHASE_L8)
    check(R, "CR-08 (c) rehacer un plan Completado cuando la unidad tiene otro proceso abierto -> 400 'otro proceso correctivo abierto'",
          st == 400 and "otro proceso correctivo abierto" in body.replace('"', ""), f"status={st} {body[:140]}")
    check(R, "CR-08 (c) nada cambia: el completado sigue 6/9/2 y el abierto 3/4/1",
          plan(ctxc["plan"]) == antes[0] and plan(pid_c2) == antes[1],
          f"completado={plan(ctxc['plan'])} abierto={plan(pid_c2)}")

    # ================================================================ CR-09 contadores del dashboard correctivo
    st_wiz, _ = c.get(f"/?ShowWizard=true&ManagementId={corr}&Step=0")
    check(R, "CR-09 el resumen del asistente correctivo (Step=0) responde 200", st_wiz == 200, f"status={st_wiz}")

    st, html = c.get(f"/Managements/Details?id={corr}&ActiveTab=dashboard")
    # Espejo SQL de Pages/Managements/Details.cshtml.cs: pipeline = CurrentState<>2 y no completados (CurrentState=9 o PlanStatus=2)
    filas = sql("select CurrentPhase, COUNT(*) from ManagementPlans where ManagementId=? and EquipmentUnitId is not null "
                "and CurrentState<>2 and not (CurrentState=9 or PlanStatus=2) group by CurrentPhase", corr)
    por_fase = {int(f): int(n) for f, n in filas}
    tiles = {lbl.strip(): int(v) for v, lbl in
             re.findall(r'asset-metric-value font-weight-bold">(\d+)</div>\s*<div class="font-weight-bold mt-2">([^<]+)</div>', html or "")}
    ok_contadores = st == 200 and all(tiles.get(lbl) == por_fase.get(fase, 0)
                                      for lbl, fase in (("L-7 Solicitud", PHASE_L7), ("L-8 Mantenimiento", PHASE_L8),
                                                        ("L-3 Salida", PHASE_L3), ("Desembolso", PHASE_L12)))
    check(R, "CR-09 contadores L-7/L-8/L-3/Desembolso del dashboard de la correctiva coinciden con SQL (planes por fase, pipeline activo)",
          ok_contadores, f"html={tiles} sql={por_fase} status={st}")
    check(R, "CR-09 CountL6 de una gestion correctiva es 0: el dashboard ni siquiera muestra la tarjeta 'L-6 Verificacion'",
          st == 200 and "L-6 Verificaci\u00f3n" not in (html or ""), f"status={st} tiles={sorted(tiles)}")

    # ================================================================ CR-10 informes en correctivo (plan llevado a DONE)
    p10 = plan(ctxc["plan"])
    err_report(R, "CR-10 L-6 en gestion correctiva -> 400 'L-6 solo aplica al proceso preventivo.'",
               c, f"/api/reports/download/l6?managementId={corr}&labId={lab}", 400, "L-6 solo aplica al proceso preventivo")
    ok_xlsx(R, "CR-10 L-7 solicitud tecnica (xlsx valido con datos)",
            c, f"/api/reports/download/l7?requestId={p10['req']}&managementId={corr}")
    ok_xlsx(R, "CR-10 L-8 kardex de mantenimiento (xlsx valido con datos)",
            c, f"/api/reports/download/l8?managementPlanId={ctxc['plan']}")
    ok_xlsx(R, "CR-10 L-3 salida (xlsx valido con datos)",
            c, f"/api/reports/download/l3?departureId={p10['dep']}")
    ok_xlsx(R, "CR-10 L-12 adquisicion (xlsx valido con datos)",
            c, f"/api/reports/download/adquisicion?requestId={p10['acq']}")
    ok_xlsx(R, "CR-10 L-48 plan de mantenimiento correctivo (xlsx valido con datos)",
            c, f"/api/reports/download/l48?managementId={corr}&labId={lab}")

    # ================================================================ CR-11 permisos: Supervisor y Cierre de practicas
    csup = P.client("qa_sup", pw)
    st, loc = csup.probe("/Requests/Index?handler=NewRequest")
    check(R, "CR-11 Supervisor: 'Nueva solicitud' (/Requests/Index?handler=NewRequest) termina en el asistente correctivo (paso 2)",
          st == 302 and "showwizard=true" in (loc or "").lower() and f"ManagementId={corr}" in (loc or "") and "Step=2" in (loc or ""),
          f"status={st} loc={loc}")
    u11, _ = P.nueva_unidad(c, lab, "c11")
    r = C.l7_corr(csup, corr, lab, u11, name="CR-11 L-7 correctiva como Supervisor", follow=False)
    pid11 = C.plan_de(corr, u11)
    check(R, "CR-11 Supervisor: POST L-7 correctiva crea el plan (302 al paso 3)", r.status == 302 and bool(pid11),
          f"{P.dump(r)} plan={plan(pid11)}")
    st, loc = csup.probe(f"/Maintenances/Create?managementPlanId={pid11}&isWizard=true")
    check(R, "CR-11 Supervisor: /Maintenances/Create del plan -> acceso denegado (302 a /Error|/Login|AccessDenied o 403, sin formulario)",
          st in (302, 303, 403) and (st == 403 or any(x in (loc or "") for x in ("/Error", "/Login", "AccessDenied"))),
          f"status={st} loc={loc}")
    cci = P.client("qa_cierre", pw)
    st, loc = cci.probe("/Requests/Create")
    check(R, "CR-11 Cierre de practicas: GET /Requests/Create redirige a /Cierre (middleware PracticeClosureAccess)",
          st == 302 and (loc or "").rstrip("/").lower().endswith("/cierre"), f"status={st} loc={loc}")

    # ================================================================ CR-12 concurrencia: dos L-7 simultaneos (D-7)
    u12, inv12 = P.nueva_unidad(c, lab, "c12")
    estados, barrera = {}, threading.Barrier(2)

    def hilo(k):
        cl = P.client("qa_super", pw)
        barrera.wait()
        r = C.l7_corr(cl, corr, lab, u12, name=f"CR-12 L-7 simultaneo hilo {k}", follow=False)
        estados[k] = r.status

    hilos = [threading.Thread(target=hilo, args=(k,)) for k in (1, 2)]
    for h in hilos:
        h.start()
    for h in hilos:
        h.join()
    n_req = scalar("select count(*) from Requests where EquipmentUnitId=? and Type=1", u12)
    n_plan = scalar("select count(*) from ManagementPlans where ManagementId=? and EquipmentUnitId=?", corr, u12)
    check(R, "CR-12 dos L-7 correctivos simultaneos para la MISMA unidad: como maximo 1 Request tecnica y 1 ManagementPlan (D-7 + indice unico)",
          n_req <= 1 and n_plan <= 1, f"unidad={inv12} requests={n_req} planes={n_plan} respuestas={estados}")
    check(R, "CR-12 ninguna de las dos respuestas es 500", all(s < 500 for s in estados.values()), f"respuestas={estados}")

    # ================================================================ CR-13 gestion correctiva NO activa
    inactiva = scalar("select top 1 Id from Managements where Type=1 and Status in (1,2) order by Status, Id")
    if not inactiva:
        P.post_form(c, "/Managements/Create?type=Corrective",
                    {"Input.Year": "2026", "Input.Semester": "1", "Input.Status": "1", "Input.Type": "Corrective",
                     "Input.StartDate": "2026-01-01", "Input.PlannedEndDate": "2026-06-30",
                     "Input.Description": "Correctiva QA inactiva", "Input.FacultyId": ""},
                    name="CR-13 crear CORR-2026-1 inactiva")
        inactiva = scalar("select top 1 Id from Managements where Code='CORR-2026-1' and Type=1 and Status=1")
    check(R, "CR-13 existe una gestion correctiva NO activa para la prueba (Status 1 o 2; se crea CORR-2026-1 si faltaba y NO se activa)",
          bool(inactiva) and scalar("select Status from Managements where Id=?", inactiva) in (1, 2),
          f"id={inactiva} status={scalar('select Status from Managements where Id=?', inactiva)}")
    u13, _ = P.nueva_unidad(c, lab, "c13")
    n_antes = scalar("select count(*) from Requests")
    campos13 = [("Input.FacultyId", str(P.FAC)), ("Input.LaboratoryId", str(lab)), ("Input.EquipmentUnitId", str(u13)),
                ("Input.Description", "Falla correctiva QAP"), ("Input.Observations", C.checklist(13)),
                ("Input.TechnicalChecklistConfirmed", "true"), ("Input.Priority", "1")]
    r = P.post_raw(c, f"/Requests/Create?isWizard=true&ManagementId={inactiva}", campos13,
                   f"/Requests/Create?isWizard=true&ManagementId={corr}", follow=True, name="CR-13 L-7 en gestion inactiva")
    check(R, "CR-13 L-7 contra una gestion correctiva NO activa -> warning 'no esta activa' y COUNT de Requests sin cambio",
          r.has_toast("no est\u00e1 activa", "warning") and scalar("select count(*) from Requests") == n_antes
          and C.plan_de(inactiva, u13) is None,
          P.dump(r) + f" requests={n_antes}->{scalar('select count(*) from Requests')}")

    # ================================================================ CR-15 convivencia: cronograma y edicion correctiva
    st_c, st_p = scalar("select Status from Managements where Id=?", corr), scalar("select Status from Managements where Id=?", P.MGMT)
    check(R, "CR-15 precondicion: la preventiva 30 y la CORR-2027-2 siguen Activas a la vez (SQL)", st_c == 0 and st_p == 0,
          f"corr={st_c} preventiva30={st_p}")
    st, html = c.get("/Schedules/Index")
    n30 = scalar("select count(*) from ManagementPlans where ManagementId=?", P.MGMT)
    m = re.search(r'<strong class="text-dark">(\d+)</strong>\s+intervenci', html or "")
    check(R, "CR-15 /Schedules/Index (preventiva por defecto): 200 y el numero de intervenciones corresponde a la gestion 30",
          st == 200 and m and int(m.group(1)) == n30, f"status={st} html={m.group(1) if m else None} sql={n30}")
    st, html = c.get("/Schedules/Index?TypeFilter=Corrective")
    ncorr = scalar("select count(*) from ManagementPlans where ManagementId=?", corr)
    m = re.search(r'<strong class="text-dark">(\d+)</strong>\s+intervenci', html or "")
    sel = re.search(r'<option[^>]*value="Corrective"[^>]*>', html or "")
    check(R, "CR-15 /Schedules/Index?TypeFilter=Corrective: 200, filtro 'Gestion correctiva' seleccionado y conteo = planes de CORR-2027-2",
          st == 200 and sel is not None and "selected" in sel.group(0) and m and int(m.group(1)) == ncorr,
          f"status={st} filtro={sel.group(0) if sel else None} html={m.group(1) if m else None} sql={ncorr}")
    req15 = plan(ctxa["plan"])["req"]
    st, html = c.get(f"/Requests/Edit/{req15}")  # ruta @page "{id:int}"
    check(R, "CR-15 /Requests/Edit de una solicitud correctiva responde 200 (sin marcador visible de correctivo en el HTML)",
          st == 200, f"status={st} codigo_gestion_visible={'CORR-2027-2' in (html or '')}")


if __name__ == "__main__":
    P.run_module("correctivo2", run)
