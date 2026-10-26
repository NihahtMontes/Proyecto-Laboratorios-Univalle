"""
Utilidades compartidas del importador único V12 -> DB_Laboratorios_Univalle_REAL.
Fuente: C:\\Users\\monte\\Downloads\\BaseGastronomia_v1.xlsx (SHA fijo). Las hojas "_*" no se importan.

Regla central: cada fila cargada deja una fila en ImportSourceRows con OriginalDataJson = la fila completa
del Excel (todas las columnas), de modo que ninguna celda se pierde aunque no tenga columna propia.
Ejecutar el SQL generado con:
  sqlcmd -S localhost -E -d DB_Laboratorios_Univalle_REAL -b -f 65001 -i <archivo.sql>
"""
import hashlib
import json
import os
import sys
from datetime import date, datetime

import openpyxl

XLSX = r"C:\Users\monte\Downloads\BaseGastronomia_v1.xlsx"
EXPECTED_SHA = "edca5f8b4a6601091a4e85939f9cf1aaa586a4e27c2e4df0dc1278afbb47a3f2"
SOURCE_NAME = "BaseGastronomia_v1.xlsx"
DB_NAME = "DB_Laboratorios_Univalle_REAL"
BATCH_CODE = "XLSX-V12-FULL-" + EXPECTED_SHA[:12].upper()
OUT_DIR = os.path.join(os.environ.get("TEMP", "."), "v12full")

# Filas esperadas por hoja (línea base del Excel V12).
EXPECTED_ROWS = {
    "Sedes": 3, "Facultades": 2, "Carreras": 4, "Laboratorios": 24, "Gestiones": 27, "Actores": 57,
    "ActorRoles": 46, "CatalogoEquipos": 121, "UnidadesFisicas": 644, "VerificacionesL6": 1599,
    "SolicitudesL7": 329, "SolicitudUnidades": 413, "MantenimientosL8": 1209, "MantSolicitudes": 121,
    "MantParticipantes": 1209, "PlanL48": 387, "Materias": 41, "Practicas": 196, "Planificacion": 422,
    "CategoriasArticulo": 10, "UnidadesMedida": 5, "Articulos": 1374, "RequerimientosPractica": 2111,
    "ExistenciasArticulos": 700, "SolicitudesL3": 1037, "SolicitudL3Detalles": 11451, "Incidentes": 0,
}

BATCH_SIZE = 100


def check_sha():
    sha = hashlib.sha256(open(XLSX, "rb").read()).hexdigest()
    if sha != EXPECTED_SHA:
        sys.exit(f"Hash distinto: {sha}")
    return sha


def read_sheets(names):
    """Devuelve {hoja: (encabezados, filas)}; cada fila es dict con '_row' = número real de fila en Excel."""
    wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)
    out = {}
    for name in names:
        if name not in wb.sheetnames:
            sys.exit(f"Falta la hoja {name}")
        it = wb[name].iter_rows(values_only=True)
        header = [h for h in next(it)]
        if any(h is None or str(h).strip() == "" for h in header):
            sys.exit(f"Hoja {name}: encabezado vacío {header}")
        rows = []
        for n, r in enumerate(wb[name].iter_rows(min_row=2, values_only=True), start=2):
            if not any(v is not None and str(v).strip() != "" for v in r):
                continue
            d = {h: (r[i] if i < len(r) else None) for i, h in enumerate(header)}
            d["_row"] = n
            rows.append(d)
        if len(rows) != EXPECTED_ROWS[name]:
            sys.exit(f"Hoja {name}: {len(rows)} filas, se esperaban {EXPECTED_ROWS[name]}")
        out[name] = (header, rows)
    return out


def text(v):
    return None if v is None or str(v).strip() == "" else str(v).strip()


def q(v):
    """Literal T-SQL; textos como N'' sin cambiar mayúsculas/minúsculas (solo se recortan los bordes)."""
    if v is None or (isinstance(v, str) and v.strip() == ""):
        return "NULL"
    if isinstance(v, bool):
        return "1" if v else "0"
    if isinstance(v, (int, float)):
        return str(int(v)) if isinstance(v, float) and v.is_integer() else repr(v)
    if isinstance(v, datetime):
        return "'" + v.strftime("%Y-%m-%dT%H:%M:%S") + "'"
    if isinstance(v, date):
        return "'" + v.strftime("%Y-%m-%d") + "'"
    return "N'" + str(v).strip().replace("'", "''") + "'"


def _json_default(v):
    if isinstance(v, (datetime, date)):
        return v.isoformat()
    return str(v)


def row_json(header, row):
    return json.dumps({h: row.get(h) for h in header}, ensure_ascii=False, default=_json_default)


def check_len(sheet, row, column, value, limit):
    if value is not None and len(str(value).strip()) > limit:
        sys.exit(f"{sheet} fila {row['_row']}: {column} excede {limit} caracteres")


class SqlBuilder:
    """Acumula T-SQL en lotes de BATCH_SIZE sentencias dentro de una transacción."""

    def __init__(self):
        self.lines = []
        self._pending = 0
        self.issue_count = 0
        self.source_row_count = 0
        self._in_group = False

    def w(self, s):
        self.lines.append(s)

    def begin_group(self):
        self.w("BEGIN TRANSACTION;")
        self._pending = 0
        self._in_group = True

    def stmt(self, s):
        self.w(s)
        if not self._in_group:
            return
        self._pending += 1
        if self._pending >= BATCH_SIZE:
            # Cada lote de 100 sentencias es una transacción y un batch aparte (GO): compila rápido y las variables se redeclaran.
            self.w("COMMIT TRANSACTION;")
            self.w("GO")
            self.w(self._declarations())
            self.w("BEGIN TRANSACTION;")
            self._pending = 0

    def end_group(self):
        self.w("COMMIT TRANSACTION;")
        self._pending = 0
        self._in_group = False

    @staticmethod
    def _declarations():
        return ("DECLARE @Now datetime2 = SYSUTCDATETIME(); "
                f"DECLARE @Batch int = (SELECT Id FROM dbo.ImportBatches WHERE Code = N'{BATCH_CODE}');")

    def prologue(self, sha, notes):
        self.w("SET NOCOUNT ON;")
        self.w("SET QUOTED_IDENTIFIER ON;")
        self.w("SET ANSI_NULLS ON; SET ANSI_PADDING ON; SET ANSI_WARNINGS ON; SET CONCAT_NULL_YIELDS_NULL ON; SET ARITHABORT ON;")
        self.w("SET XACT_ABORT ON;")
        self.w(f"IF DB_NAME() <> N'{DB_NAME}' THROW 50000, 'BD incorrecta: solo {DB_NAME}', 1;")
        self.w("DECLARE @Now datetime2 = SYSUTCDATETIME();")
        self.w(f"IF NOT EXISTS (SELECT 1 FROM dbo.ImportBatches WHERE Code = N'{BATCH_CODE}') "
               f"INSERT dbo.ImportBatches(Code, SourceType, SourceName, SourceSha256, ImportedAt, Status, Notes, ContractVersion) "
               f"VALUES(N'{BATCH_CODE}', N'Excel', N'{SOURCE_NAME}', N'{sha}', @Now, 1, {q(notes)}, N'base-gastronomia-v12-full');")
        self.w(f"DECLARE @Batch int = (SELECT Id FROM dbo.ImportBatches WHERE Code = N'{BATCH_CODE}');")

    def source_row(self, sheet, header, row, key, entity, target_key):
        """Registra la fila completa del Excel (red de seguridad: nada se pierde)."""
        k = f"{sheet}:{key}"
        self.source_row_count += 1
        self.stmt(
            "IF NOT EXISTS (SELECT 1 FROM dbo.ImportSourceRows WHERE ImportBatchId = @Batch AND "
            f"SourceRowKey = {q(k)} AND TargetEntityName = {q(entity)}) "
            "INSERT dbo.ImportSourceRows(ImportBatchId, SourceSheet, SourceRowNumber, SourceRowKey, OriginalIdentifier, "
            "TargetEntityName, TargetEntityKey, MigrationStatus, ReconciliationStatus, OriginalDataJson, CreatedDate) "
            f"VALUES(@Batch, {q(sheet)}, {row['_row']}, {q(k)}, {q(str(key))}, {q(entity)}, {q(str(target_key))}, 1, 2, "
            f"{q(row_json(header, row))}, @Now);")

    def issue(self, code, entity, key, field, original, normalized, desc, severity, sheet, row):
        """Incidencia de calidad idempotente por (lote, código, entidad, clave, campo)."""
        self.issue_count += 1
        field_cmp = q(field) if field else "N''"
        self.stmt(
            "IF NOT EXISTS (SELECT 1 FROM dbo.DataQualityIssues WHERE ImportBatchId = @Batch AND "
            f"IssueCode = {q(code)} AND EntityName = {q(entity)} AND EntityKey = {q(str(key))} AND "
            f"ISNULL(FieldName, N'') = {field_cmp}) "
            "INSERT dbo.DataQualityIssues(ImportBatchId, IssueCode, EntityName, EntityKey, FieldName, OriginalValue, "
            "NormalizedValue, Description, Severity, Status, CreatedDate, SourceSheet, SourceRowNumber, SourceCode) "
            f"VALUES(@Batch, {q(code)}, {q(entity)}, {q(str(key))}, {q(field)}, {q(original)}, {q(normalized)}, "
            f"{q(desc)}, {severity}, 0, @Now, {q(sheet)}, {row['_row'] if row else 'NULL'}, {q(str(key))});")

    def write(self, name):
        os.makedirs(OUT_DIR, exist_ok=True)
        path = os.path.join(OUT_DIR, name)
        with open(path, "w", encoding="utf-8-sig", newline="\n") as f:
            f.write("\n".join(self.lines) + "\n")
        return path
