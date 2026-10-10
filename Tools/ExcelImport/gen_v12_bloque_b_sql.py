"""
V12 · Bloque B de BaseGastronomia_v1.xlsx: circuito L-3 real de 2026-1 y 2026-2.
Genera import_v12_bloque_b.sql SOLO para DB_Laboratorios_Univalle_DEMO_QA.
Ejecutar el SQL con: sqlcmd -S localhost -E -C -d DB_Laboratorios_Univalle_DEMO_QA -b -f 65001 -i <sql>

Contrato: docs/INSTRUCCIONES_AGENTE_V12_GASTRONOMIA.md (§4 y decisiones D1, D2, D6, D7, D8).
- Calendario de 2026-1 y 2026-2: fechas de clases = primera y última práctica de la fuente (derivadas, no
  inventadas); semanas lunes-sábado con la misma regla que AcademicCalendarService.GenerateWeeksAsync.
- CourseGroups derivados GRP-<gestión>-<MATERIA>-<GRUPO>, PracticeSchedules (D8) para las solicitudes que
  resuelven a una práctica del catálogo, Plannings (SolicitudesL3) y PlanningRequirements (SolicitudL3Detalles).
- D1: RequestDate es NOT NULL y la fuente no la trae → RequestDate = fecha de práctica + una incidencia de lote.
- D2: las solicitudes sin fecha o sin laboratorio no se cargan (ni sus líneas).
- D6: horarios que no calzan con los periodos quedan sin periodo y con la hora de la fuente.
- Debe ejecutarse ANTES de la migración AddStudentGroups: su SQL convierte estos grupos y renombra GroupName.
- No crea salidas, devoluciones ni incidentes, y no toca nada histórico.
"""
import collections
import hashlib
import io
import os
import sys
from datetime import date, datetime, time, timedelta

import openpyxl

XLSX = r"C:\Users\monte\Downloads\BaseGastronomia_v1.xlsx"
EXPECTED_SHA = "edca5f8b4a6601091a4e85939f9cf1aaa586a4e27c2e4df0dc1278afbb47a3f2"
SOURCE_NAME = "BaseGastronomia_v1.xlsx"
OUT = os.path.join(os.environ.get("TEMP", "."), "mm", "import_v12_bloque_b.sql")

HEADERS = {
    "SolicitudesL3": ("SOLICITUDL3CODIGO", "GESTIONCODIGO", "SEDECODIGO", "CARRERACODIGO", "LABORATORIOCODIGO",
                      "TODOSLOSLABORATORIOS", "DOCENTEACTORCODIGO", "MATERIACODIGO", "GRUPO", "FECHAPRACTICA",
                      "HORAINICIO", "HORAFIN", "NUMEROPRACTICA", "ESTADO", "OBSERVACIONES"),
    "SolicitudL3Detalles": ("SOLICITUDL3DETALLECODIGO", "SOLICITUDL3CODIGO", "ARTICULOCODIGO", "NUMEROINVENTARIO",
                            "CANTIDADSOLICITADA", "OBSERVACIONES"),
}
# Catálogo fijo de ClassPeriods (Id = Order); el SQL aborta si la BD no coincide.
PERIODS = [(1, "07:35", "08:25"), (2, "08:35", "09:25"), (3, "09:25", "10:15"), (4, "10:25", "11:15"),
           (5, "11:15", "12:05"), (6, "12:15", "13:05"), (7, "13:05", "13:55"), (8, "14:05", "14:55"),
           (9, "14:55", "15:45"), (10, "15:55", "16:45"), (11, "16:45", "17:35"), (12, "17:45", "18:35")]
STATUS = {"aprobada": ("Aprobada", 0), "terminada": ("Finalizada", 1), "parcial": ("Finalizada", 1),
          "cancelada": ("Cancelada", 3)}  # → (Planning.Status, PracticeScheduleStatus)

sha = hashlib.sha256(open(XLSX, "rb").read()).hexdigest()
if sha != EXPECTED_SHA:
    sys.exit(f"Hash distinto: {sha}")


def read_book(path, names):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out = {}
    for name in names:
        it = wb[name].iter_rows(values_only=True)
        header = next(it)
        seen = {(h or "").strip().casefold(): i for i, h in enumerate(header) if (h or "").strip()}
        mapping = {}
        for c in HEADERS[name]:
            if c.casefold() not in seen:
                sys.exit(f"Hoja {name}: falta encabezado '{c}'")
            mapping[c] = seen[c.casefold()]
        rows = []
        for n, r in enumerate(wb[name].iter_rows(min_row=2, values_only=True), start=2):
            if not any(v is not None and str(v).strip() != "" for v in r):
                continue
            d = {c: (r[i] if i < len(r) else None) for c, i in mapping.items()}
            d["_row"] = n
            rows.append(d)
        out[name] = rows
    return out


def code(v):
    return "" if v is None else str(v).strip().upper()


def text(v):
    return None if v is None or str(v).strip() == "" else str(v).strip()


def q(v):
    if v is None or (isinstance(v, str) and v.strip() == ""):
        return "NULL"
    if isinstance(v, bool):
        return "1" if v else "0"
    if isinstance(v, (int, float)):
        return str(int(v)) if isinstance(v, float) and v.is_integer() else str(v)
    if isinstance(v, datetime):
        return "'" + v.strftime("%Y-%m-%d") + "'"
    if isinstance(v, date):
        return "'" + v.strftime("%Y-%m-%d") + "'"
    if isinstance(v, time):
        return "'" + v.strftime("%H:%M") + "'"
    return "N'" + str(v).strip().replace("'", "''") + "'"


def hm(t):
    return None if t is None else t.strftime("%H:%M")


B = read_book(XLSX, list(HEADERS))
reqs = B["SolicitudesL3"]
lines = B["SolicitudL3Detalles"]
for rows in B.values():
    for r in rows:
        for c in list(r):
            if c.endswith("CODIGO") or c in ("GRUPO", "NUMEROPRACTICA"):
                r[c] = code(r[c]) if r[c] is not None else None

# ---------- Pre-vuelo ----------
assert len(reqs) == 1037, len(reqs)
assert len(lines) == 11451, len(lines)
assert len({r["SOLICITUDL3CODIGO"] for r in reqs}) == 1037, "Códigos de solicitud repetidos"
assert len({l["SOLICITUDL3DETALLECODIGO"] for l in lines}) == 11451, "Códigos de línea repetidos"
req_codes = {r["SOLICITUDL3CODIGO"] for r in reqs}
assert all(l["SOLICITUDL3CODIGO"] in req_codes for l in lines), "Línea con solicitud inexistente"
assert all(not text(l["NUMEROINVENTARIO"]) for l in lines), "Hay números de inventario: revisar mapeo"
assert all(r["TODOSLOSLABORATORIOS"] in (False, None, "No") for r in reqs)
assert {str(r["ESTADO"]).strip().lower() for r in reqs} <= set(STATUS), {r["ESTADO"] for r in reqs}
assert all(r["GESTIONCODIGO"] in ("2026-1", "2026-2") for r in reqs)

excluded = [r for r in reqs if r["FECHAPRACTICA"] is None or not r["LABORATORIOCODIGO"]]
assert len(excluded) == 9, [r["SOLICITUDL3CODIGO"] for r in excluded]
excluded_codes = {r["SOLICITUDL3CODIGO"] for r in excluded}
loaded = [r for r in reqs if r["SOLICITUDL3CODIGO"] not in excluded_codes]
loaded_lines = [l for l in lines if l["SOLICITUDL3CODIGO"] not in excluded_codes]
assert all(r["FECHAPRACTICA"].weekday() != 6 for r in loaded), "Práctica en domingo"

# D6: periodos
starts = {s: pid for pid, s, _ in PERIODS}
ends = {e: pid for pid, _, e in PERIODS}
for r in loaded:
    s, e = hm(r["HORAINICIO"]), hm(r["HORAFIN"])
    r["_p"] = None
    if s in starts and e in ends and ends[e] >= starts[s]:
        r["_p"] = (starts[s], ends[e] - starts[s] + 1)
    if r["HORAINICIO"] and r["HORAFIN"] and r["HORAFIN"] <= r["HORAINICIO"]:
        sys.exit(f"Hora de fin no posterior al inicio: {r['SOLICITUDL3CODIGO']}")
aligned = sum(1 for r in loaded if r["_p"])

# D7: fechas de clases derivadas de la fuente
class_range = {}
for m in ("2026-1", "2026-2"):
    ds = [r["FECHAPRACTICA"].date() for r in loaded if r["GESTIONCODIGO"] == m]
    class_range[m] = (min(ds), max(ds))


def week_slots(start, end):
    monday = start - timedelta(days=start.weekday())
    last = end - timedelta(days=end.weekday())
    n = 1
    while monday <= last:
        yield n, monday
        n += 1
        monday += timedelta(days=7)


# Grupos derivados: docente y laboratorio más frecuentes
combos = collections.defaultdict(list)
for r in loaded:
    combos[(r["GESTIONCODIGO"], r["MATERIACODIGO"], r["GRUPO"])].append(r)
assert all(len(g) <= 20 for _, _, g in combos), "Nombre de grupo > 20"

# Choques de horario en la fuente (informativos)
overlaps = []
by_day = collections.defaultdict(list)
for r in loaded:
    if r["HORAINICIO"] and r["HORAFIN"]:
        by_day[r["FECHAPRACTICA"].date()].append(r)
for day_rows in by_day.values():
    day_rows.sort(key=lambda x: x["SOLICITUDL3CODIGO"])
    for i, a in enumerate(day_rows):
        for b in day_rows[i + 1:]:
            if a["HORAINICIO"] < b["HORAFIN"] and b["HORAINICIO"] < a["HORAFIN"]:
                if a["LABORATORIOCODIGO"] == b["LABORATORIOCODIGO"]:
                    overlaps.append((a, b, "Laboratoriocodigo"))
                if a["DOCENTEACTORCODIGO"] == b["DOCENTEACTORCODIGO"]:
                    overlaps.append((a, b, "Docenteactorcodigo"))

# Líneas repetidas (solicitud, artículo)
pair_count = collections.Counter((l["SOLICITUDL3CODIGO"], l["ARTICULOCODIGO"]) for l in loaded_lines)

print(f"[pre] solicitudes={len(reqs)} cargables={len(loaded)} excluidas={len(excluded)} "
      f"lineas_cargables={len(loaded_lines)} con_periodo={aligned} grupos={len(combos)} "
      f"choques_fuente={len(overlaps)} pares_repetidos={sum(1 for v in pair_count.values() if v > 1)} "
      f"clases={ {k: (str(a), str(b)) for k, (a, b) in class_range.items()} }", flush=True)

# ---------- SQL ----------
sql = []
w = sql.append
batch_code = "XLSX-V12-L3-" + sha[:12].upper()
w("SET NOCOUNT ON;")
w("SET QUOTED_IDENTIFIER ON; SET ANSI_NULLS ON; SET ANSI_PADDING ON; SET ANSI_WARNINGS ON; "
  "SET CONCAT_NULL_YIELDS_NULL ON; SET ARITHABORT ON;")
w("SET XACT_ABORT ON;")
w("IF DB_NAME() <> N'DB_Laboratorios_Univalle_DEMO_QA' THROW 50000, 'Destino incorrecto: solo DEMO_QA', 1;")
w("IF NOT EXISTS (SELECT 1 FROM dbo.__EFMigrationsHistory WHERE MigrationId = N'20261009141121_RestructurePlanningModel') "
  "THROW 50001, 'Falta la migración RestructurePlanningModel.', 1;")
w("IF OBJECT_ID(N'dbo.StudentGroups') IS NOT NULL THROW 50002, "
  "'AddStudentGroups ya está aplicada: este bloque debe correr antes para que su conversión alcance a los grupos.', 1;")
w("IF (SELECT COUNT(*) FROM dbo.ClassPeriods) <> 12 OR EXISTS (SELECT 1 FROM dbo.ClassPeriods WHERE Id <> [Order]) "
  "THROW 50003, 'El catálogo de periodos no es el esperado (12, Id = Order).', 1;")
for pid, s, e in PERIODS:
    w(f"IF NOT EXISTS (SELECT 1 FROM dbo.ClassPeriods WHERE Id = {pid} AND StartTime = '{s}' AND EndTime = '{e}') "
      f"THROW 50004, 'Periodo {pid} distinto al esperado.', 1;")
w("DECLARE @Now datetime2 = SYSUTCDATETIME();")
w("DECLARE @n int;")
w(f"IF NOT EXISTS (SELECT 1 FROM dbo.ImportBatches WHERE Code = N'{batch_code}') "
  f"INSERT dbo.ImportBatches(Code, SourceType, SourceName, SourceSha256, ImportedAt, Status, Notes, ContractVersion) "
  f"VALUES(N'{batch_code}', N'Excel', N'{SOURCE_NAME}', N'{sha}', @Now, 1, "
  f"N'V12 bloque B: SolicitudesL3/SolicitudL3Detalles 2026-1 y 2026-2 → Plannings/PlanningRequirements. Solo DEMO_QA.', "
  f"N'base-gastronomia-v12');")
w(f"DECLARE @Batch int = (SELECT Id FROM dbo.ImportBatches WHERE Code = N'{batch_code}');")
w("CREATE TABLE #Issue(IssueCode nvarchar(100), EntityName nvarchar(100), EntityKey nvarchar(200), FieldName nvarchar(100), "
  "OriginalValue nvarchar(400), NormalizedValue nvarchar(400), Description nvarchar(1000), Severity int, SourceSheet nvarchar(128), SourceRow int);")


def issue(icode, entity, key, field, original, normalized, desc, severity, sheet_name, row):
    w(f"INSERT #Issue VALUES({q(icode)}, {q(entity)}, {q(key)}, {q(field)}, {q(original)}, {q(normalized)}, "
      f"{q(desc)}, {severity}, {q(sheet_name)}, {row if row is not None else 'NULL'});")


# 1. Calendario académico (D7)
w("-- 1. Calendario de 2026-1 y 2026-2: fechas de clases derivadas + semanas lunes-sábado (regla de AcademicCalendarService)")
w("CREATE TABLE #Slot(Mgmt nvarchar(50), Number int, Monday date, Saturday date, PRIMARY KEY(Mgmt, Monday));")
w("CREATE TABLE #Day(Mgmt nvarchar(50), Monday date, [Date] date, PRIMARY KEY(Mgmt, [Date]));")
w("CREATE TABLE #Range(Mgmt nvarchar(50) PRIMARY KEY, FromDate date, ToDate date);")
for m, (a, b) in class_range.items():
    w(f"INSERT #Range VALUES(N'{m}', {q(a)}, {q(b)});")
    for n, monday in week_slots(a, b):
        w(f"INSERT #Slot VALUES(N'{m}', {n}, {q(monday)}, {q(monday + timedelta(days=5))});")
        for i in range(6):
            d = monday + timedelta(days=i)
            if a <= d <= b:
                w(f"INSERT #Day VALUES(N'{m}', {q(monday)}, {q(d)});")
    issue("VALUE_DERIVED_FROM_SOURCE", "Management", m, "ClassesStartDate/ClassesEndDate", None, f"{a} / {b}",
          "La fuente no trae el inicio y fin de clases: se toman la primera y la última práctica L-3 de la gestión. "
          "Confirmar con la clienta (semana 1 = semana de la primera práctica).", 1, "SolicitudesL3", None)
w("""IF EXISTS (SELECT 1 FROM dbo.AcademicWeeks w JOIN dbo.Managements m ON m.Id = w.ManagementId JOIN #Range r ON r.Mgmt = m.Code
           WHERE NOT EXISTS (SELECT 1 FROM #Slot s WHERE s.Mgmt = m.Code AND s.Monday = w.StartDate)
             AND (EXISTS (SELECT 1 FROM dbo.Plannings p WHERE p.AcademicWeekId = w.Id)
                  OR EXISTS (SELECT 1 FROM dbo.PracticeSchedules p WHERE p.AcademicWeekId = w.Id)
                  OR EXISTS (SELECT 1 FROM dbo.AcademicDays d WHERE d.AcademicWeekId = w.Id AND d.DayRecoveryId IS NOT NULL)))
    THROW 50010, 'Una semana fuera del nuevo rango tiene planificaciones o reposiciones: revisar a mano.', 1;
IF EXISTS (SELECT 1 FROM dbo.AcademicDays d JOIN dbo.AcademicWeeks w ON w.Id = d.AcademicWeekId
           JOIN dbo.Managements m ON m.Id = w.ManagementId JOIN #Range r ON r.Mgmt = m.Code
           WHERE (d.[Date] < r.FromDate OR d.[Date] > r.ToDate)
             AND (d.DayRecoveryId IS NOT NULL OR EXISTS (SELECT 1 FROM dbo.Plannings p WHERE p.ManagementId = m.Id
                  AND p.IsDeleted = 0 AND CAST(p.ScheduledDate AS date) = d.[Date])))
    THROW 50011, 'Un día fuera del nuevo rango tiene planificaciones o reposiciones: revisar a mano.', 1;
BEGIN TRANSACTION;
UPDATE m SET ClassesStartDate = r.FromDate, ClassesEndDate = r.ToDate, LastModifiedDate = @Now
FROM dbo.Managements m JOIN #Range r ON r.Mgmt = m.Code
WHERE m.ClassesStartDate IS NULL AND m.ClassesEndDate IS NULL;
SELECT 'gestiones_fechas_clases' k, @@ROWCOUNT n;
DELETE d FROM dbo.AcademicDays d JOIN dbo.AcademicWeeks w ON w.Id = d.AcademicWeekId
JOIN dbo.Managements m ON m.Id = w.ManagementId JOIN #Range r ON r.Mgmt = m.Code
WHERE d.[Date] < r.FromDate OR d.[Date] > r.ToDate
   OR NOT EXISTS (SELECT 1 FROM #Slot s WHERE s.Mgmt = m.Code AND s.Monday = w.StartDate);
SELECT 'dias_fuera_de_rango_quitados' k, @@ROWCOUNT n;
DELETE w FROM dbo.AcademicWeeks w JOIN dbo.Managements m ON m.Id = w.ManagementId JOIN #Range r ON r.Mgmt = m.Code
WHERE NOT EXISTS (SELECT 1 FROM #Slot s WHERE s.Mgmt = m.Code AND s.Monday = w.StartDate);
SELECT 'semanas_fuera_de_rango_quitadas' k, @@ROWCOUNT n;
CREATE TABLE #Renum(WeekId int PRIMARY KEY, NewNumber int);
INSERT #Renum SELECT w.Id, s.Number FROM dbo.AcademicWeeks w JOIN dbo.Managements m ON m.Id = w.ManagementId
JOIN #Slot s ON s.Mgmt = m.Code AND s.Monday = w.StartDate WHERE w.Number <> s.Number;
UPDATE w SET Number = -w.Number FROM dbo.AcademicWeeks w JOIN #Renum r ON r.WeekId = w.Id;
UPDATE w SET Number = r.NewNumber FROM dbo.AcademicWeeks w JOIN #Renum r ON r.WeekId = w.Id;
SELECT 'semanas_renumeradas' k, COUNT(*) n FROM #Renum;
INSERT dbo.AcademicWeeks(Number, StartDate, EndDate, ManagementId)
SELECT s.Number, s.Monday, s.Saturday, m.Id FROM #Slot s JOIN dbo.Managements m ON m.Code = s.Mgmt
WHERE NOT EXISTS (SELECT 1 FROM dbo.AcademicWeeks w WHERE w.ManagementId = m.Id AND w.StartDate = s.Monday);
SELECT 'semanas_insertadas' k, @@ROWCOUNT n;
INSERT dbo.AcademicDays(AcademicWeekId, [Date], Status)
SELECT w.Id, d.[Date], 0 FROM #Day d JOIN dbo.Managements m ON m.Code = d.Mgmt
JOIN dbo.AcademicWeeks w ON w.ManagementId = m.Id AND w.StartDate = d.Monday
WHERE NOT EXISTS (SELECT 1 FROM dbo.AcademicDays x WHERE x.AcademicWeekId = w.Id AND x.[Date] = d.[Date]);
SELECT 'dias_insertados' k, @@ROWCOUNT n;
UPDATE p SET AcademicWeek = w.Number FROM dbo.Plannings p JOIN dbo.AcademicWeeks w ON w.Id = p.AcademicWeekId
JOIN #Renum r ON r.WeekId = w.Id;
UPDATE p SET WeekNumber = w.Number FROM dbo.PracticeSchedules p JOIN dbo.AcademicWeeks w ON w.Id = p.AcademicWeekId
JOIN #Renum r ON r.WeekId = w.Id;
COMMIT TRANSACTION;""")

# 2. Solicitudes a tabla temporal
w("-- 2. Solicitudes cargables")
w("CREATE TABLE #Req(Code nvarchar(50) PRIMARY KEY, Mgmt nvarchar(50), Campus nvarchar(20), Career nvarchar(20), Lab nvarchar(50), "
  "Teacher nvarchar(50), Subj nvarchar(50), Grp nvarchar(20), PDate date, T1 time NULL, T2 time NULL, Num nvarchar(20), "
  "PStatus nvarchar(50), SStatus int, Obs nvarchar(1000), StartPeriod int NULL, PeriodCount int NULL, SourceRow int);")
for r in loaded:
    p_status, s_status = STATUS[str(r["ESTADO"]).strip().lower()]
    sp, pc = r["_p"] if r["_p"] else (None, None)
    w(f"INSERT #Req VALUES({q(r['SOLICITUDL3CODIGO'])}, {q(r['GESTIONCODIGO'])}, {q(r['SEDECODIGO'])}, "
      f"{q(r['CARRERACODIGO'])}, {q(r['LABORATORIOCODIGO'])}, {q(r['DOCENTEACTORCODIGO'])}, {q(r['MATERIACODIGO'])}, "
      f"{q(r['GRUPO'])}, {q(r['FECHAPRACTICA'])}, {q(r['HORAINICIO'])}, {q(r['HORAFIN'])}, {q(r['NUMEROPRACTICA'])}, "
      f"N'{p_status}', {s_status}, {q(text(r['OBSERVACIONES']))}, {q(sp)}, {q(pc)}, {r['_row']});")
    if not r["_p"]:
        issue("TIME_NOT_ALIGNED_TO_PERIOD", "Planning", r["SOLICITUDL3CODIGO"], "Horainicio/Horafin",
              f"{hm(r['HORAINICIO']) or '-'} / {hm(r['HORAFIN']) or '-'}", None,
              "El horario de la fuente no coincide con los límites de los periodos: se carga sin periodo y con la hora "
              "de la fuente (D6). No se redondea.", 1, "SolicitudesL3", r["_row"])
for r in excluded:
    issue("REQUIRED_VALUE_MISSING", "Planning", r["SOLICITUDL3CODIGO"], "Fechapractica/Laboratoriocodigo",
          f"{q(r['FECHAPRACTICA']) if r['FECHAPRACTICA'] else '-'} / {r['LABORATORIOCODIGO'] or '-'}", None,
          "La solicitud no tiene fecha de práctica o laboratorio: no se carga, ni sus líneas (D2). Pedir el dato a la clienta.",
          2, "SolicitudesL3", r["_row"])
for a, b, field in overlaps:
    issue("SCHEDULE_OVERLAP_SOURCE", "Planning", f"{a['SOLICITUDL3CODIGO']}|{b['SOLICITUDL3CODIGO']}", field,
          a[field.upper()], None,
          "Dos solicitudes de la fuente se solapan el mismo día (mismo laboratorio o docente). Informativo: se cargan ambas.",
          0, "SolicitudesL3", a["_row"])
issue("VALUE_PENDING_CONFIRMATION", "Planning", "SolicitudesL3", "Fechasolicitud", None, "RequestDate = Fechapractica",
      "La fuente no trae la fecha de solicitud en ninguna fila y RequestDate es obligatorio: se usa la fecha de práctica (D1).",
      1, "SolicitudesL3", None)

w("""INSERT #Issue SELECT 'UNRESOLVED_REFERENCE', 'Planning', r.Code, 'Gestion/Sede/Lab/Docente/Materia',
   CONCAT(r.Mgmt, N' / ', r.Campus, N' / ', r.Lab, N' / ', r.Teacher, N' / ', r.Subj), NULL,
   N'Referencia inexistente; la solicitud no se carga.', 2, 'SolicitudesL3', r.SourceRow
FROM #Req r WHERE NOT EXISTS (SELECT 1 FROM dbo.Managements m WHERE m.Code = r.Mgmt)
            OR NOT EXISTS (SELECT 1 FROM dbo.Campuses c WHERE c.Code = r.Campus)
            OR NOT EXISTS (SELECT 1 FROM dbo.Laboratories l WHERE l.Code = r.Lab)
            OR NOT EXISTS (SELECT 1 FROM dbo.People p WHERE p.ActorCode = r.Teacher)
            OR NOT EXISTS (SELECT 1 FROM dbo.Subjects s WHERE s.Code = r.Subj);
INSERT #Issue SELECT 'CAREER_MISMATCH', 'Planning', r.Code, 'Carreracodigo', r.Career, c.Code,
   N'La carrera del formulario no es la de la materia; se usa la de la materia (Planning no guarda carrera).', 0, 'SolicitudesL3', r.SourceRow
FROM #Req r JOIN dbo.Subjects s ON s.Code = r.Subj JOIN dbo.Careers c ON c.Id = s.CareerId WHERE c.Code <> r.Career;
INSERT #Issue SELECT 'PRACTICE_NOT_IN_CATALOG', 'Planning', r.Code, 'Numeropractica', CONCAT(r.Subj, N' / ', r.Num), NULL,
   N'La práctica no existe en el catálogo de la materia: se carga sin programación y con nombre derivado (D8).', 0, 'SolicitudesL3', r.SourceRow
FROM #Req r JOIN dbo.Subjects s ON s.Code = r.Subj
WHERE NOT EXISTS (SELECT 1 FROM dbo.Practices p WHERE p.SubjectId = s.Id AND p.Status <> 2 AND UPPER(LTRIM(RTRIM(p.Number))) = r.Num);""")

# 3. Grupos derivados
w("-- 3. Grupos derivados GRP-<gestión>-<MATERIA>-<GRUPO> (docente y laboratorio más frecuentes)")
w("CREATE TABLE #Grp(Code nvarchar(50) PRIMARY KEY, Mgmt nvarchar(50), Subj nvarchar(50), Grp nvarchar(20), Teacher nvarchar(50), Lab nvarchar(50));")
for (m, s, g), rows in combos.items():
    gcode = f"GRP-{m}-{s}-{g}"
    picked = {}
    for field in ("DOCENTEACTORCODIGO", "LABORATORIOCODIGO"):
        counts = collections.Counter(r[field] for r in rows)
        picked[field] = counts.most_common(1)[0][0]
        if len(counts) > 1:
            issue("COURSE_GROUP_VALUES_VARY", "CourseGroup", gcode, field, None, ", ".join(sorted(counts)),
                  "Las solicitudes del grupo tienen valores distintos; el grupo toma el más frecuente y cada solicitud conserva el suyo.",
                  0, "SolicitudesL3", None)
    w(f"INSERT #Grp VALUES({q(gcode)}, {q(m)}, {q(s)}, {q(g)}, {q(picked['DOCENTEACTORCODIGO'])}, {q(picked['LABORATORIOCODIGO'])});")
w("""INSERT #Issue SELECT 'CONTENT_CONFLICT', 'CourseGroup', g.Code, 'Gestion/Materia/Grupo', x.Code, NULL,
   N'Ya existe otro grupo activo con la misma gestión, materia y nombre: se reutiliza.', 0, 'SolicitudesL3', NULL
FROM #Grp g JOIN dbo.Managements m ON m.Code = g.Mgmt JOIN dbo.Subjects s ON s.Code = g.Subj
JOIN dbo.CourseGroups x ON x.ManagementId = m.Id AND x.SubjectId = s.Id AND x.Name = g.Grp AND x.Status <> 2 AND x.Code <> g.Code;
BEGIN TRANSACTION;
INSERT dbo.CourseGroups(Code, ManagementId, SubjectId, Name, TeacherPersonId, LaboratoryId, IsDemonstrative, Status, CreatedDate)
SELECT g.Code, m.Id, s.Id, g.Grp, p.Id, l.Id, 0, 0, @Now
FROM #Grp g JOIN dbo.Managements m ON m.Code = g.Mgmt JOIN dbo.Subjects s ON s.Code = g.Subj
JOIN dbo.People p ON p.ActorCode = g.Teacher LEFT JOIN dbo.Laboratories l ON l.Code = g.Lab
WHERE NOT EXISTS (SELECT 1 FROM dbo.CourseGroups x WHERE x.Code = g.Code)
  AND NOT EXISTS (SELECT 1 FROM dbo.CourseGroups x WHERE x.ManagementId = m.Id AND x.SubjectId = s.Id AND x.Name = g.Grp AND x.Status <> 2);
SELECT 'grupos_insertados' k, @@ROWCOUNT n;
COMMIT TRANSACTION;""")

# 4. Programaciones (D8) para las solicitudes que resuelven a una práctica
w("""-- 4. PracticeSchedules (D8): una por solicitud que resuelve a una práctica del catálogo (código PROG-<SOL3>)
INSERT #Issue SELECT 'CONTENT_CONFLICT', 'PracticeSchedule', CONCAT(N'PROG-', r.Code), 'Gestion/Practica/Grupo', z.Code, NULL,
   N'Ya existe otra programación con la misma gestión, práctica y grupo: la solicitud se carga sin programación.', 1, 'SolicitudesL3', r.SourceRow
FROM #Req r JOIN dbo.Managements m ON m.Code = r.Mgmt JOIN dbo.Subjects s ON s.Code = r.Subj
JOIN dbo.Practices pr ON pr.SubjectId = s.Id AND pr.Status <> 2 AND UPPER(LTRIM(RTRIM(pr.Number))) = r.Num
JOIN dbo.PracticeSchedules z ON z.ManagementId = m.Id AND z.PracticeId = pr.Id AND z.GroupName = r.Grp AND z.IsDeleted = 0
                            AND z.Code <> CONCAT(N'PROG-', r.Code);
SET @n = 1;
WHILE @n > 0
BEGIN
    BEGIN TRANSACTION;
    INSERT dbo.PracticeSchedules(Code, ManagementId, SubjectId, PracticeId, CourseGroupId, GroupName, TeacherPersonId,
        LaboratoryId, WeekNumber, AcademicWeekId, ExecutionDate, StartTime, EndTime, Status, ImportBatchId, CreatedDate, IsDeleted)
    SELECT TOP (100) CONCAT(N'PROG-', r.Code), m.Id, s.Id, pr.Id, cg.Id, r.Grp, pe.Id, l.Id, w.Number, w.Id,
        r.PDate, r.T1, r.T2, r.SStatus, @Batch, @Now, 0
    FROM #Req r JOIN dbo.Managements m ON m.Code = r.Mgmt JOIN dbo.Subjects s ON s.Code = r.Subj
    JOIN dbo.Practices pr ON pr.SubjectId = s.Id AND pr.Status <> 2 AND UPPER(LTRIM(RTRIM(pr.Number))) = r.Num
    JOIN dbo.People pe ON pe.ActorCode = r.Teacher JOIN dbo.Laboratories l ON l.Code = r.Lab
    LEFT JOIN dbo.AcademicWeeks w ON w.ManagementId = m.Id AND r.PDate BETWEEN w.StartDate AND w.EndDate
    LEFT JOIN dbo.CourseGroups cg ON cg.ManagementId = m.Id AND cg.SubjectId = s.Id AND cg.Name = r.Grp AND cg.Status <> 2
    WHERE NOT EXISTS (SELECT 1 FROM dbo.PracticeSchedules z WHERE z.Code = CONCAT(N'PROG-', r.Code))
      AND NOT EXISTS (SELECT 1 FROM dbo.PracticeSchedules z
                      WHERE z.ManagementId = m.Id AND z.PracticeId = pr.Id AND z.GroupName = r.Grp AND z.IsDeleted = 0);
    SET @n = @@ROWCOUNT;
    COMMIT TRANSACTION;
END;
SELECT 'programaciones_l3' k, COUNT(*) n FROM dbo.PracticeSchedules WHERE Code LIKE N'PROG-SOL3-%';""")

# 5. Plannings
w("""-- 5. Plannings (SolicitudesL3). RequestDate = fecha de práctica (D1); periodo solo si calza exacto (D6).
SET @n = 1;
WHILE @n > 0
BEGIN
    BEGIN TRANSACTION;
    INSERT dbo.Plannings(Code, ManagementId, CampusId, LaboratoryId, AllLaboratories, TeacherPersonId, SubjectId,
        CourseGroupId, GroupName, PracticeNumber, PracticeName, PracticeScheduleId, AcademicWeekId, AcademicWeek,
        ScheduledDate, ScheduledEndDate, StartTime, EndTime, StartPeriodId, PeriodCount, Status, RequestDate,
        Observations, CreatedDate, IsDeleted)
    SELECT TOP (100) r.Code, m.Id, c.Id, l.Id, 0, pe.Id, s.Id, cg.Id, r.Grp, r.Num,
        LEFT(COALESCE(pr.Name, CONCAT(N'Práctica ', r.Num, N' – ', s.Name)), 255),
        ps.Id, w.Id, ISNULL(w.Number, 0),
        DATEADD(minute, ISNULL(DATEDIFF(minute, CAST('00:00' AS time), r.T1), 0), CAST(r.PDate AS datetime2)),
        CASE WHEN r.T2 IS NULL THEN NULL ELSE DATEADD(minute, DATEDIFF(minute, CAST('00:00' AS time), r.T2), CAST(r.PDate AS datetime2)) END,
        r.T1, r.T2, r.StartPeriod, r.PeriodCount, r.PStatus, CAST(r.PDate AS datetime2), r.Obs, @Now, 0
    FROM #Req r JOIN dbo.Managements m ON m.Code = r.Mgmt JOIN dbo.Campuses c ON c.Code = r.Campus
    JOIN dbo.Laboratories l ON l.Code = r.Lab JOIN dbo.People pe ON pe.ActorCode = r.Teacher
    JOIN dbo.Subjects s ON s.Code = r.Subj
    LEFT JOIN dbo.Practices pr ON pr.SubjectId = s.Id AND pr.Status <> 2 AND UPPER(LTRIM(RTRIM(pr.Number))) = r.Num
    LEFT JOIN dbo.PracticeSchedules ps ON ps.Code = CONCAT(N'PROG-', r.Code)
    LEFT JOIN dbo.AcademicWeeks w ON w.ManagementId = m.Id AND r.PDate BETWEEN w.StartDate AND w.EndDate
    LEFT JOIN dbo.CourseGroups cg ON cg.ManagementId = m.Id AND cg.SubjectId = s.Id AND cg.Name = r.Grp AND cg.Status <> 2
    WHERE NOT EXISTS (SELECT 1 FROM dbo.Plannings x WHERE x.Code = r.Code);
    SET @n = @@ROWCOUNT;
    COMMIT TRANSACTION;
END;
SELECT 'plannings_l3' k, COUNT(*) n FROM dbo.Plannings WHERE Code LIKE N'SOL3-%';""")

# 6. Líneas
w("-- 6. PlanningRequirements (SolicitudL3Detalles)")
w("CREATE TABLE #Line(Code nvarchar(50) PRIMARY KEY, Req nvarchar(50), Art nvarchar(30), Qty decimal(18,2) NULL, Obs nvarchar(1000), SourceRow int);")
for l in loaded_lines:
    qty = l["CANTIDADSOLICITADA"]
    lc = l["SOLICITUDL3DETALLECODIGO"]
    w(f"INSERT #Line VALUES({q(lc)}, {q(l['SOLICITUDL3CODIGO'])}, {q(l['ARTICULOCODIGO'])}, "
      f"{q(qty) if qty not in (None, 0) else 'NULL'}, {q(text(l['OBSERVACIONES']))}, {l['_row']});")
    if qty is None:
        issue("VALUE_MISSING", "PlanningRequirement", lc, "Cantidadsolicitada", None, None,
              "La línea no trae cantidad: se carga con cantidad vacía.", 1, "SolicitudL3Detalles", l["_row"])
    elif qty == 0:
        issue("VALUE_ZERO_IN_SOURCE", "PlanningRequirement", lc, "Cantidadsolicitada", "0", None,
              "Cantidad 0 en la fuente: se carga vacía (la BD exige NULL o > 0).", 1, "SolicitudL3Detalles", l["_row"])
    elif qty >= 100:
        issue("QUANTITY_UNIT_SUSPECT", "PlanningRequirement", lc, "Cantidadsolicitada", str(qty), None,
              "Cantidad ≥ 100 en un artículo por pieza: probablemente gramos. Se carga tal cual; confirmar unidad.",
              0, "SolicitudL3Detalles", l["_row"])
    if pair_count[(l["SOLICITUDL3CODIGO"], l["ARTICULOCODIGO"])] > 1:
        issue("DUPLICATE_LINE_IN_SOURCE", "PlanningRequirement", lc, "Articulocodigo", l["ARTICULOCODIGO"],
              l["SOLICITUDL3CODIGO"], "El artículo se repite en la misma solicitud: se cargan ambas líneas, sin sumar.",
              0, "SolicitudL3Detalles", l["_row"])
w("""INSERT #Issue SELECT 'UNRESOLVED_REFERENCE', 'PlanningRequirement', x.Code, 'Articulocodigo', x.Art, NULL,
   N'Artículo inexistente o dado de baja; la línea no se carga.', 2, 'SolicitudL3Detalles', x.SourceRow
FROM #Line x WHERE NOT EXISTS (SELECT 1 FROM dbo.Articles a WHERE a.Code = x.Art AND a.Status <> 2);
SET @n = 1;
WHILE @n > 0
BEGIN
    BEGIN TRANSACTION;
    INSERT dbo.PlanningRequirements(Code, PlanningId, ArticleId, EquipmentUnitId, RequestedQuantity, ExpectedQuantity,
        DeliveredQuantity, ReturnedQuantity, Observations)
    SELECT TOP (100) x.Code, p.Id, a.Id, NULL, x.Qty, CAST(CEILING(ISNULL(x.Qty, 0)) AS int), 0, 0, x.Obs
    FROM #Line x JOIN dbo.Plannings p ON p.Code = x.Req JOIN dbo.Articles a ON a.Code = x.Art AND a.Status <> 2
    WHERE NOT EXISTS (SELECT 1 FROM dbo.PlanningRequirements z WHERE z.Code = x.Code);
    SET @n = @@ROWCOUNT;
    COMMIT TRANSACTION;
END;
SELECT 'lineas_l3' k, COUNT(*) n FROM dbo.PlanningRequirements z JOIN dbo.Plannings p ON p.Id = z.PlanningId WHERE p.Code LIKE N'SOL3-%';""")

# 7. Incidencias
w("""-- 7. Incidencias (idempotentes por ImportBatchId + IssueCode + EntityKey + FieldName)
SET @n = 1;
WHILE @n > 0
BEGIN
    BEGIN TRANSACTION;
    INSERT dbo.DataQualityIssues(ImportBatchId, IssueCode, EntityName, EntityKey, FieldName, OriginalValue, NormalizedValue,
        Description, Severity, Status, CreatedDate, SourceSheet, SourceRowNumber)
    SELECT TOP (100) @Batch, i.IssueCode, i.EntityName, i.EntityKey, i.FieldName, i.OriginalValue, i.NormalizedValue,
        i.Description, i.Severity, 0, @Now, i.SourceSheet, i.SourceRow
    FROM (SELECT DISTINCT * FROM #Issue) i
    WHERE NOT EXISTS (SELECT 1 FROM dbo.DataQualityIssues d WHERE d.ImportBatchId = @Batch AND d.IssueCode = i.IssueCode
                      AND ISNULL(d.EntityKey, N'') = ISNULL(i.EntityKey, N'')
                      AND ISNULL(d.FieldName, N'') = ISNULL(i.FieldName, N''));
    SET @n = @@ROWCOUNT;
    COMMIT TRANSACTION;
END;
SELECT IssueCode, COUNT(*) n FROM dbo.DataQualityIssues WHERE ImportBatchId = @Batch GROUP BY IssueCode ORDER BY IssueCode;""")

# 8. Verificación
w("""-- 8. Verificación
SELECT m.Code gestion, COUNT(*) plannings, SUM(CASE WHEN p.StartPeriodId IS NOT NULL THEN 1 ELSE 0 END) con_periodo,
       SUM(CASE WHEN p.PracticeScheduleId IS NOT NULL THEN 1 ELSE 0 END) con_programacion,
       SUM(CASE WHEN p.AcademicWeekId IS NULL THEN 1 ELSE 0 END) sin_semana
FROM dbo.Plannings p JOIN dbo.Managements m ON m.Id = p.ManagementId WHERE p.Code LIKE N'SOL3-%' GROUP BY m.Code;
SELECT m.Code gestion, COUNT(*) semanas, MIN(w.StartDate) desde, MAX(w.EndDate) hasta
FROM dbo.AcademicWeeks w JOIN dbo.Managements m ON m.Id = w.ManagementId WHERE m.Code IN (N'2026-1', N'2026-2') GROUP BY m.Code;
SELECT 'grupos_l3' k, COUNT(*) n FROM dbo.CourseGroups WHERE Code LIKE N'GRP-2026-%';""")

os.makedirs(os.path.dirname(OUT), exist_ok=True)
io.open(OUT, "w", encoding="utf-8").write("\n".join(sql) + "\n")
print(OUT, f"| lote={batch_code}", flush=True)
