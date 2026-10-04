"""Genera import_v3_demo.sql desde BaseEstructuraFinal_V3_ARTICULOS.xlsx (solo DB_Laboratorios_Univalle_DEMO_QA)."""
import collections
import hashlib
import io
import os
import unicodedata

import openpyxl

XLSX = r"C:\Users\monte\Downloads\BaseEstructuraFinal_V3_ARTICULOS.xlsx"
EXPECTED_SHA = "1e70f5e4cbf9b298ee3690bfc65e4a0aec07efae109ca398aa10b774ca87680e"
OUT = os.path.join(os.environ.get("TEMP", "."), "mm", "import_v3_demo.sql")

sha = hashlib.sha256(open(XLSX, "rb").read()).hexdigest()
assert sha == EXPECTED_SHA, f"Hash distinto: {sha}"

wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)


def sheet(name):
    it = wb[name].iter_rows(values_only=True)
    header = next(it)
    out = []
    for i, r in enumerate(it, start=2):
        if any(v is not None for v in r):
            d = dict(zip(header, r))
            d["_row"] = i
            out.append(d)
    return out


def q(v):
    if v is None or (isinstance(v, str) and v.strip() == ""):
        return "NULL"
    return "N'" + str(v).strip().replace("'", "''") + "'"


def norm_words(s):
    s = unicodedata.normalize("NFKD", str(s)).encode("ascii", "ignore").decode().upper()
    stop = {"DE", "CON", "Y", "EL", "LA", "LOS", "LAS", "PARA", "EN", "X"}
    words = "".join(ch if ch.isalnum() else " " for ch in s).split()
    # Mismas palabras en otro orden, ignorando plural simple y conectores (criterio del informe: 21 grupos).
    return frozenset((w[:-1] if w.endswith("S") and len(w) > 3 else w) for w in words if w not in stop)


cats = sheet("CategoriasArticulo")
units = sheet("UnidadesMedida")
arts = sheet("Articulos")
actors = sheet("Actores")
roles = sheet("ActorRoles")
cat_name = {c["CategoriaCodigo"]: c["Nombre"] for c in cats}
unit_name = {u["UnidadMedidaCodigo"]: u["Nombre"] for u in units}
assert len(cats) == 10 and len(units) == 3 and len(arts) == 1023, (len(cats), len(units), len(arts))

# Alias "Nombre en la planificación" -> ActorCodigo (bitácora _POBLADO_DOCENTES).
aliases = {}
for r in wb["_POBLADO_DOCENTES"].iter_rows(values_only=True):
    if r and isinstance(r[1], str) and r[1].startswith("ACT-00") and isinstance(r[0], str):
        aliases[r[1]] = r[0].strip()

NEW_ACTORS = {f"ACT-{n:04d}" for n in list(range(37, 52)) + [8, 9] + list(range(27, 37))}
actor_rows = [a for a in actors if a["ActorCodigo"] in NEW_ACTORS]
assert len(actor_rows) == 27, len(actor_rows)

ROLE_MAP = {"técnico": 1, "tecnico": 1, "proveedor": 3, "solicitante": 5, "responsable": 6, "verificador": 7, "docente": 8}
roles_by_actor = collections.defaultdict(set)
for r in roles:
    if r["ActorCodigo"] in NEW_ACTORS and str(r.get("Estado") or "Activo").lower() == "activo":
        roles_by_actor[r["ActorCodigo"]].add(ROLE_MAP.get(str(r["Rol"]).strip().lower(), 99))

# Casi-duplicados: mismas palabras en otro orden.
groups = collections.defaultdict(list)
for a in arts:
    groups[norm_words(a["Nombre"])].append(a["ArticuloCodigo"])
dup_groups = [codes for key, codes in groups.items() if len(codes) > 1]
SUSPICIOUS = {"ECO-R0349": "Nombre posiblemente truncado (102 → 92 caracteres)",
              "ART-REQ-00492": "Nombre incompleto",
              "ART-REQ-00479": "Posible errata ('REDONDOBORDE')"}

sql = []
w = sql.append
w("SET NOCOUNT ON;")
w("SET XACT_ABORT ON;")
w("IF DB_NAME() <> N'DB_Laboratorios_Univalle_DEMO_QA' THROW 50000, 'Destino incorrecto: solo DEMO_QA', 1;")
w("DECLARE @Now datetime2 = SYSUTCDATETIME();")
w("DECLARE @Ins TABLE(Kind nvarchar(30), N int);")
w("")
w("-- 0. Lote de importación V3")
batch_code = "XLSX-V3-ARTICULOS-" + sha[:12].upper()
w(f"IF NOT EXISTS (SELECT 1 FROM dbo.ImportBatches WHERE Code = N'{batch_code}') "
  f"INSERT dbo.ImportBatches(Code, SourceType, SourceName, SourceSha256, ImportedAt, Status, Notes, ContractVersion) "
  f"VALUES(N'{batch_code}', N'Excel', N'BaseEstructuraFinal_V3_ARTICULOS.xlsx', N'{sha}', @Now, 1, "
  f"N'Artículos, categorías, unidades y docentes (V3). Solo DEMO_QA.', N'base-estructural-v3');")
w(f"DECLARE @Batch int = (SELECT Id FROM dbo.ImportBatches WHERE Code = N'{batch_code}');")
w("")
w("-- 1. Categorías y unidades por código. Si un registro demo tiene el mismo nombre, adopta el código de la V3.")
w("BEGIN TRANSACTION;")
for c in cats:
    code, name = q(c["CategoriaCodigo"]), q(c["Nombre"])
    w(f"UPDATE dbo.ArticleCategories SET Code = {code} WHERE Name = {name} AND Code <> {code} AND NOT EXISTS (SELECT 1 FROM dbo.ArticleCategories x WHERE x.Code = {code});")
    w(f"IF NOT EXISTS (SELECT 1 FROM dbo.ArticleCategories WHERE Code = {code}) BEGIN INSERT dbo.ArticleCategories(Code, Name, Status, CreatedDate) VALUES({code}, {name}, 0, @Now); INSERT @Ins VALUES('categoria', 1); END")
    w(f"UPDATE dbo.ArticleCategories SET Name = {name}, Status = 0 WHERE Code = {code} AND (Name <> {name} OR Status <> 0);")
for u in units:
    code, name = q(u["UnidadMedidaCodigo"]), q(u["Nombre"])
    w(f"UPDATE dbo.MeasurementUnits SET Code = {code} WHERE Name = {name} AND Code <> {code} AND NOT EXISTS (SELECT 1 FROM dbo.MeasurementUnits x WHERE x.Code = {code});")
    w(f"IF NOT EXISTS (SELECT 1 FROM dbo.MeasurementUnits WHERE Code = {code}) BEGIN INSERT dbo.MeasurementUnits(Code, Name, Status, CreatedDate) VALUES({code}, {name}, 0, @Now); INSERT @Ins VALUES('unidad', 1); END")
    w(f"UPDATE dbo.MeasurementUnits SET Name = {name}, Status = 0 WHERE Code = {code} AND (Name <> {name} OR Status <> 0);")
w("COMMIT TRANSACTION;")
w("")
w("-- 2. Artículos (lotes de 100, una transacción por lote)")
w("CREATE TABLE #V3Art(Code nvarchar(30) NOT NULL PRIMARY KEY, Name nvarchar(200) NOT NULL, CatCode nvarchar(50) NOT NULL, UnitCode nvarchar(50) NOT NULL, Pending bit NOT NULL, SourceRow int NOT NULL);")
for i in range(0, len(arts), 100):
    chunk = arts[i:i + 100]
    w("INSERT #V3Art VALUES " + ", ".join(
        f"({q(a['ArticuloCodigo'])}, {q(a['Nombre'])}, {q(a['CategoriaCodigo'])}, {q(a['UnidadMedidaCodigo'])}, "
        f"{1 if str(a['Estado']).strip().lower() == 'por confirmar' else 0}, {a['_row']})" for a in chunk) + ";")
w("""
-- 2a. Artículos existentes (con otro código) cuyo nombre choca con uno de la V3: baja lógica y traslado de receta y stock.
CREATE TABLE #Clash(OldId int NOT NULL, OldCode nvarchar(30) NOT NULL, NewCode nvarchar(30) NOT NULL);
INSERT #Clash SELECT a.Id, a.Code, v.Code FROM dbo.Articles a JOIN #V3Art v ON v.Name = a.Name AND v.Code <> a.Code WHERE a.Status <> 2;
BEGIN TRANSACTION;
UPDATE a SET Status = 2, LastModifiedDate = @Now FROM dbo.Articles a JOIN #Clash c ON c.OldId = a.Id;
COMMIT TRANSACTION;

DECLARE @n int = 1;
WHILE @n > 0
BEGIN
    BEGIN TRANSACTION;
    INSERT dbo.Articles(Code, Name, Category, UnitOfMeasure, Status, CreatedDate, ImportBatchId, ArticleCategoryId, MeasurementUnitId, ReviewStatus)
    SELECT TOP (100) v.Code, v.Name, c.Name, u.Name, 0, @Now, @Batch, c.Id, u.Id, CASE WHEN v.Pending = 1 THEN 9 ELSE 0 END
    FROM #V3Art v
    JOIN dbo.ArticleCategories c ON c.Code = v.CatCode
    JOIN dbo.MeasurementUnits u ON u.Code = v.UnitCode
    WHERE NOT EXISTS (SELECT 1 FROM dbo.Articles a WHERE a.Code = v.Code);
    SET @n = @@ROWCOUNT;
    INSERT @Ins VALUES('articulo', @n);
    COMMIT TRANSACTION;
END;

-- Actualización idempotente de artículos V3 ya existentes (lotes de 100).
SET @n = 1;
WHILE @n > 0
BEGIN
    BEGIN TRANSACTION;
    UPDATE TOP (100) a SET Name = v.Name, Category = c.Name, UnitOfMeasure = u.Name, ArticleCategoryId = c.Id, MeasurementUnitId = u.Id,
        ReviewStatus = CASE WHEN v.Pending = 1 THEN 9 ELSE 0 END, LastModifiedDate = @Now
    FROM dbo.Articles a JOIN #V3Art v ON v.Code = a.Code
    JOIN dbo.ArticleCategories c ON c.Code = v.CatCode JOIN dbo.MeasurementUnits u ON u.Code = v.UnitCode
    WHERE a.Name <> v.Name OR a.ArticleCategoryId <> c.Id OR a.MeasurementUnitId <> u.Id
       OR a.ReviewStatus <> CASE WHEN v.Pending = 1 THEN 9 ELSE 0 END OR a.Category <> c.Name OR a.UnitOfMeasure <> u.Name;
    SET @n = @@ROWCOUNT;
    COMMIT TRANSACTION;
END;

-- 2b. Traslado de receta y stock de los artículos dados de baja al artículo V3 equivalente.
BEGIN TRANSACTION;
UPDATE r SET ArticleId = na.Id, LastModifiedDate = @Now
FROM dbo.PracticeRequirements r
JOIN #Clash c ON c.OldId = r.ArticleId
JOIN dbo.Articles na ON na.Code = c.NewCode
WHERE r.IsRemoved = 0
  AND NOT EXISTS (SELECT 1 FROM dbo.PracticeRequirements x WHERE x.PracticeId = r.PracticeId AND x.ArticleId = na.Id AND x.IsRemoved = 0);
UPDATE s SET ArticleId = na.Id, LastModifiedDate = @Now
FROM dbo.ArticleStocks s
JOIN #Clash c ON c.OldId = s.ArticleId
JOIN dbo.Articles na ON na.Code = c.NewCode
WHERE NOT EXISTS (SELECT 1 FROM dbo.ArticleStocks x WHERE x.LaboratoryId = s.LaboratoryId AND x.ArticleId = na.Id);
COMMIT TRANSACTION;
""")
w("-- 3. Actores nuevos (15 docentes V3 + 12 faltantes de la V1)")
w("BEGIN TRANSACTION;")
for a in actor_rows:
    code = a["ActorCodigo"]
    link = str(a.get("Vinculacion") or "").strip().lower()
    is_entity = str(a.get("TipoActor") or "").strip().lower() == "entidad"
    actor_roles = roles_by_actor.get(code, set())
    if 8 in actor_roles:
        category = 2
    elif 1 in actor_roles:
        category = 1
    elif link == "externo" or 3 in actor_roles:
        category = 5
    else:
        category = 99
    status = 0 if str(a.get("Estado") or "Activo").strip().lower() == "activo" else 1
    w(f"IF NOT EXISTS (SELECT 1 FROM dbo.People WHERE ActorCode = N'{code}') BEGIN "
      f"INSERT dbo.People(Status, Category, Email, PhoneNumber, CreatedDate, ImportBatchId, ActorCode, Title) "
      f"VALUES({status}, {category}, {q(a.get('Email'))}, {q(a.get('Telefono'))}, @Now, @Batch, N'{code}', {q(a.get('Titulo'))}); INSERT @Ins VALUES('actor', 1); END")
    pid = f"(SELECT Id FROM dbo.People WHERE ActorCode = N'{code}')"
    if link == "externo" or is_entity:
        w(f"IF NOT EXISTS (SELECT 1 FROM dbo.Externs WHERE Id = {pid}) AND NOT EXISTS (SELECT 1 FROM dbo.Interns WHERE Id = {pid}) "
          f"INSERT dbo.Externs(Id, IsEntity, Name, Address, ExternStatus) VALUES({pid}, {1 if is_entity else 0}, {q(a['NombreRazonSocial'])}, {q(a.get('Direccion')) if a.get('Direccion') else "N''"}, {status});")
    else:
        w(f"IF NOT EXISTS (SELECT 1 FROM dbo.Interns WHERE Id = {pid}) AND NOT EXISTS (SELECT 1 FROM dbo.Externs WHERE Id = {pid}) "
          f"INSERT dbo.Interns(Id, Name, InternStatus) VALUES({pid}, {q(a['NombreRazonSocial'])}, {status});")
    for role in sorted(actor_roles):
        w(f"IF NOT EXISTS (SELECT 1 FROM dbo.PersonRoleAssignments WHERE PersonId = {pid} AND Role = {role} AND IsActive = 1) "
          f"INSERT dbo.PersonRoleAssignments(PersonId, Role, IsActive, ValidFrom) VALUES({pid}, {role}, 1, @Now);")
    names = [(a["NombreRazonSocial"], 1)]
    if code in aliases and aliases[code].upper() != str(a["NombreRazonSocial"]).upper():
        names.append((aliases[code], 0))
    for alias, preferred in names:
        normalized = " ".join(str(alias).upper().split())
        w(f"IF NOT EXISTS (SELECT 1 FROM dbo.PersonAliases WHERE PersonId = {pid} AND NormalizedAlias = {q(normalized)}) "
          f"INSERT dbo.PersonAliases(PersonId, Alias, NormalizedAlias, IsPreferred, Source) VALUES({pid}, {q(alias)}, {q(normalized)}, {preferred}, N'BaseEstructuraFinal_V3_ARTICULOS.xlsx');")
w("COMMIT TRANSACTION;")
w("")
w("-- 4. Incidencias de calidad del lote V3 (idempotentes por lote + código + clave)")
w("BEGIN TRANSACTION;")
w("INSERT dbo.DataQualityIssues(ImportBatchId, IssueCode, EntityName, EntityKey, FieldName, OriginalValue, Description, Severity, Status, CreatedDate, SourceSheet, SourceRowNumber)")
w("SELECT @Batch, N'ARTICLE_PENDING_CONFIRMATION', N'Article', v.Code, N'Estado', N'Por confirmar', N'La fuente marca el artículo como \"Por confirmar\": se importa Activo con ReviewStatus = Por confirmar.', 1, 0, @Now, N'Articulos', v.SourceRow")
w("FROM #V3Art v WHERE v.Pending = 1 AND NOT EXISTS (SELECT 1 FROM dbo.DataQualityIssues d WHERE d.ImportBatchId = @Batch AND d.IssueCode = N'ARTICLE_PENDING_CONFIRMATION' AND d.EntityKey = v.Code);")
for a in actor_rows:
    if str(a.get("Vinculacion") or "").strip().lower() == "por confirmar":
        w(f"IF NOT EXISTS (SELECT 1 FROM dbo.DataQualityIssues WHERE ImportBatchId = @Batch AND IssueCode = N'ACTOR_LINK_PENDING' AND EntityKey = N'{a['ActorCodigo']}') "
          f"INSERT dbo.DataQualityIssues(ImportBatchId, IssueCode, EntityName, EntityKey, FieldName, OriginalValue, Description, Severity, Status, CreatedDate, SourceSheet, SourceRowNumber) "
          f"VALUES(@Batch, N'ACTOR_LINK_PENDING', N'Person', N'{a['ActorCodigo']}', N'Vinculacion', N'Por confirmar', N'Vinculación Interno/Externo sin confirmar; se registró como interno.', 1, 0, @Now, N'Actores', {a['_row']});")
for i, codes in enumerate(dup_groups, start=1):
    key = f"DUP-{i:02d}"
    w(f"IF NOT EXISTS (SELECT 1 FROM dbo.DataQualityIssues WHERE ImportBatchId = @Batch AND IssueCode = N'ARTICLE_POSSIBLE_DUPLICATE' AND EntityKey = {q(codes[0])}) "
      f"INSERT dbo.DataQualityIssues(ImportBatchId, IssueCode, EntityName, EntityKey, Description, CandidateKeys, Severity, Status, CreatedDate, SourceSheet) "
      f"VALUES(@Batch, N'ARTICLE_POSSIBLE_DUPLICATE', N'Article', {q(codes[0])}, N'Posible duplicado (mismas palabras en otro orden). No se fusionó.', {q(', '.join(codes))}, 1, 0, @Now, N'Articulos');")
for code, desc in SUSPICIOUS.items():
    w(f"IF NOT EXISTS (SELECT 1 FROM dbo.DataQualityIssues WHERE ImportBatchId = @Batch AND IssueCode = N'ARTICLE_SUSPICIOUS_NAME' AND EntityKey = N'{code}') "
      f"INSERT dbo.DataQualityIssues(ImportBatchId, IssueCode, EntityName, EntityKey, Description, Severity, Status, CreatedDate, SourceSheet) "
      f"VALUES(@Batch, N'ARTICLE_SUSPICIOUS_NAME', N'Article', N'{code}', {q(desc)}, 1, 0, @Now, N'Articulos');")
w("COMMIT TRANSACTION;")
w("")
w("-- Verificación")
w("SELECT Kind, SUM(N) insertados FROM @Ins GROUP BY Kind;")
w("SELECT 'categorias V3' k, COUNT(*) n FROM dbo.ArticleCategories WHERE Code LIKE N'CAT-ART-%'")
w("UNION ALL SELECT 'unidades V3', COUNT(*) FROM dbo.MeasurementUnits WHERE Code LIKE N'UNI-%'")
w("UNION ALL SELECT 'articulos V3', COUNT(*) FROM dbo.Articles a JOIN #V3Art v ON v.Code = a.Code")
w("UNION ALL SELECT 'articulos V3 por confirmar', COUNT(*) FROM dbo.Articles a JOIN #V3Art v ON v.Code = a.Code WHERE a.ReviewStatus = 9")
w("UNION ALL SELECT 'huerfanos FK articulo', COUNT(*) FROM dbo.Articles a JOIN #V3Art v ON v.Code = a.Code WHERE a.ArticleCategoryId IS NULL OR a.MeasurementUnitId IS NULL")
w("UNION ALL SELECT 'docentes con rol Docente', COUNT(DISTINCT r.PersonId) FROM dbo.PersonRoleAssignments r WHERE r.Role = 8 AND r.IsActive = 1")
w("UNION ALL SELECT 'actores con codigo', COUNT(*) FROM dbo.People WHERE ActorCode LIKE N'ACT-%'")
w("UNION ALL SELECT 'incidencias lote V3', COUNT(*) FROM dbo.DataQualityIssues WHERE ImportBatchId = @Batch;")
w("SELECT OldCode, NewCode FROM #Clash;")

io.open(OUT, "w", encoding="utf-8").write("\n".join(sql) + "\n")
print(OUT, "| articulos", len(arts), "| actores", len(actor_rows), "| alias planif", len(aliases),
      "| grupos dup", len(dup_groups), "| por confirmar", sum(1 for a in arts if str(a['Estado']).strip().lower() == 'por confirmar'))
