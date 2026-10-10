"""
V12 · Bloque A de BaseGastronomia_v1.xlsx (internamente "V12" de la clienta).
Genera import_v12_bloque_a.sql SOLO para DB_Laboratorios_Univalle_DEMO_QA.
Ejecutar el SQL con: sqlcmd -S localhost -E -C -d DB_Laboratorios_Univalle_DEMO_QA -b -f 65001 -i <sql>
(-f 65001 es obligatorio: el archivo es UTF-8 y sin esa opción sqlcmd daña las tildes).

Contrato: docs/INSTRUCCIONES_AGENTE_V12_GASTRONOMIA.md (§3 y decisiones D3, D4, D5, D7).
- Bloque A: unidad UNI-05, 22 materias, 6 docentes (+ rol, alias), 372 artículos, vida útil de
  121 equipos, existencias (agregar/actualizar sin borrar) y fechas de gestión donde son NULL.
- El Bloque B (SolicitudesL3 → Plannings) NO está aquí: depende de la migración del refactor Planning.
- No se reimporta ni se actualiza nada histórico ni ningún nombre existente.
"""
import collections
import hashlib
import io
import os
import re
import sys
from datetime import date, datetime

import openpyxl

XLSX = r"C:\Users\monte\Downloads\BaseGastronomia_v1.xlsx"
EXPECTED_SHA = "edca5f8b4a6601091a4e85939f9cf1aaa586a4e27c2e4df0dc1278afbb47a3f2"
PREV_XLSX = r"C:\Users\monte\Downloads\BaseEstructuraFinal_V7.xlsx"
SOURCE_NAME = "BaseGastronomia_v1.xlsx"
OUT = os.path.join(os.environ.get("TEMP", "."), "mm", "import_v12_bloque_a.sql")

HEADERS = {
    "Gestiones": ("GESTIONCODIGO", "TIPOGESTION", "ANIO", "SEMESTRE", "FECHAINICIO", "FECHAFIN", "ESTADO"),
    "Actores": ("ACTORCODIGO", "NOMBRERAZONSOCIAL", "TITULO", "TIPOACTOR", "VINCULACION", "EMAIL",
                "TELEFONO", "DIRECCION", "ESTADO"),
    "ActorRoles": ("ACTORROLCODIGO", "ACTORCODIGO", "ROL", "ESTADO"),
    "CatalogoEquipos": ("CATALOGOCODIGO", "VIDAUTILANIOS"),
    "Materias": ("MATERIACODIGO", "NOMBRE", "CARRERACODIGO", "SEMESTRE", "ESTADO"),
    "UnidadesMedida": ("UNIDADMEDIDACODIGO", "NOMBRE"),
    "Articulos": ("ARTICULOCODIGO", "NOMBRE", "CATEGORIACODIGO", "UNIDADMEDIDACODIGO", "ESTADO"),
    "ExistenciasArticulos": ("LABORATORIOCODIGO", "ARTICULOCODIGO", "CANTIDADDISPONIBLE", "FECHAACTUALIZACION"),
}
ROLE_MAP = {"técnico": 1, "tecnico": 1, "proveedor": 3, "solicitante": 5, "responsable": 6, "verificador": 7, "docente": 8}

sha = hashlib.sha256(open(XLSX, "rb").read()).hexdigest()
if sha != EXPECTED_SHA:
    sys.exit(f"Hash distinto: {sha}")


def read_book(path, names):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out = {}
    for name in names:
        it = wb[name].iter_rows(values_only=True)
        header = next(it)
        seen = {}
        for i, h in enumerate(header):
            key = (h or "").strip().casefold()
            if key:
                if key in seen:
                    sys.exit(f"Hoja {name}: encabezado duplicado '{h}'")
                seen[key] = i
        mapping = {}
        for c in HEADERS[name]:
            if c.casefold() not in seen:
                sys.exit(f"Hoja {name}: falta encabezado '{c}' (leídos={list(header)})")
            mapping[c] = seen[c.casefold()]
        rows = []
        for n, r in enumerate(wb[name].iter_rows(min_row=2, values_only=True), start=2):
            if not any(v is not None and str(v).strip() != "" for v in r):
                continue
            d = {c: (r[i] if i < len(r) else None) for c, i in mapping.items()}
            d["_row"] = n
            rows.append(d)
        out[name] = rows
    return out


def code(v):
    return "" if v is None else str(v).strip().upper()


def text(v):
    return None if v is None or str(v).strip() == "" else str(v).strip()


def q(v):
    if v is None or (isinstance(v, str) and v.strip() == ""):
        return "NULL"
    if isinstance(v, bool):
        return "1" if v else "0"
    if isinstance(v, (int, float)):
        return str(int(v)) if isinstance(v, float) and v.is_integer() else str(v)
    if isinstance(v, datetime):
        return "'" + v.strftime("%Y-%m-%d") + "'"
    if isinstance(v, date):
        return "'" + v.strftime("%Y-%m-%d") + "'"
    return "N'" + str(v).strip().replace("'", "''") + "'"


B = read_book(XLSX, list(HEADERS))
P = read_book(PREV_XLSX, ["Actores", "ActorRoles", "Materias", "Articulos", "ExistenciasArticulos"])

# ---------- Normalización ----------
for sheet_name, rows in B.items():
    for r in rows:
        for c in list(r):
            if c.endswith("CODIGO"):
                r[c] = code(r[c])
for sheet_name, rows in P.items():
    for r in rows:
        for c in list(r):
            if c.endswith("CODIGO"):
                r[c] = code(r[c])

prev_actors = {r["ACTORCODIGO"] for r in P["Actores"]}
prev_roles = {r["ACTORROLCODIGO"] for r in P["ActorRoles"]}
prev_subjects = {r["MATERIACODIGO"] for r in P["Materias"]}
prev_articles = {r["ARTICULOCODIGO"] for r in P["Articulos"]}

new_actors = [r for r in B["Actores"] if r["ACTORCODIGO"] not in prev_actors]
new_roles = [r for r in B["ActorRoles"] if r["ACTORROLCODIGO"] not in prev_roles]
new_subjects = [r for r in B["Materias"] if r["MATERIACODIGO"] not in prev_subjects]
new_articles = [r for r in B["Articulos"] if r["ARTICULOCODIGO"] not in prev_articles]
stocks = B["ExistenciasArticulos"]
catalog = B["CatalogoEquipos"]
all_articles = {r["ARTICULOCODIGO"] for r in B["Articulos"]}

# ---------- Pre-vuelo ----------
assert len(B["Actores"]) == 57 and len(new_actors) == 6, (len(B["Actores"]), len(new_actors))
assert len(B["ActorRoles"]) == 46 and len(new_roles) == 6, (len(B["ActorRoles"]), len(new_roles))
assert len(B["Materias"]) == 41 and len(new_subjects) == 22, (len(B["Materias"]), len(new_subjects))
assert len(B["Articulos"]) == 1374 and len(new_articles) == 372, (len(B["Articulos"]), len(new_articles))
assert not (prev_articles - all_articles), "La V12 elimina artículos de la V7"
assert len(stocks) == 700, len(stocks)
assert all(s["LABORATORIOCODIGO"] and s["CANTIDADDISPONIBLE"] is not None for s in stocks), "Existencia sin lab o cantidad"
assert not {s["ARTICULOCODIGO"] for s in stocks} - all_articles, "Existencia con artículo inexistente"
assert len({(s["LABORATORIOCODIGO"], s["ARTICULOCODIGO"]) for s in stocks}) == 700, "Existencias duplicadas"
assert len(catalog) == 121 and all(isinstance(c["VIDAUTILANIOS"], (int, float)) for c in catalog), "Vida útil incompleta"
roles_by_actor = collections.defaultdict(set)
for r in new_roles:
    role = ROLE_MAP.get(str(r["ROL"]).strip().lower())
    assert role is not None, f"Rol desconocido {r['ROL']}"
    assert r["ACTORCODIGO"] in {a["ACTORCODIGO"] for a in new_actors}, f"Rol nuevo de actor existente: {r}"
    roles_by_actor[r["ACTORCODIGO"]].add(role)
units_used = {r["UNIDADMEDIDACODIGO"] for r in new_articles}
assert units_used <= {"UNI-01", "UNI-02", "UNI-03", "UNI-05"}, units_used
for r in new_articles:
    assert len(text(r["NOMBRE"])) <= 190, r  # deja lugar para el sufijo de D4
for r in new_subjects:
    assert len(text(r["NOMBRE"])) <= 255, r

# D4: nombres repetidos dentro de los nuevos → el segundo lleva el código como sufijo.
seen_names = {}
for r in new_articles:
    r["_name"] = text(r["NOMBRE"])
    key = r["_name"].casefold()
    r["_dup_of"] = seen_names.get(key)
    if r["_dup_of"]:
        r["_name"] = f"{r['_name']} ({r['ARTICULOCODIGO']})"
    else:
        seen_names[key] = r["ARTICULOCODIGO"]
dups = [r for r in new_articles if r["_dup_of"]]
assert [(r["ARTICULOCODIGO"], r["_dup_of"]) for r in dups] == [("ART-L33-0256", "ART-L33-0007")], dups


# D3: actores en MAYÚSCULAS como los 15 docentes existentes; título separado del nombre.
TITLE_RE = re.compile(r"^(ing|lic|dra|dr|chef|tec|arq)\.\s*", re.IGNORECASE)
for a in new_actors:
    raw = text(a["NOMBRERAZONSOCIAL"])
    title = text(a["TITULO"])
    m = TITLE_RE.match(raw)
    if m:
        title = title or m.group(0).strip()
        raw = raw[m.end():].strip()
    a["_name"] = " ".join(raw.upper().split())
    a["_title"] = title.upper() if title else None
    assert a["_name"] and len(a["_name"]) <= 100, a
    assert roles_by_actor[a["ACTORCODIGO"]] == {8}, a

print(f"[pre] actores_nuevos={len(new_actors)} roles_nuevos={len(new_roles)} materias_nuevas={len(new_subjects)} "
      f"articulos_nuevos={len(new_articles)} existencias={len(stocks)} vida_util={len(catalog)} dup_nombre={len(dups)}",
      flush=True)

# ---------- SQL ----------
sql = []
w = sql.append
batch_code = "XLSX-V12-GASTRO-" + sha[:12].upper()
w("SET NOCOUNT ON;")
w("SET QUOTED_IDENTIFIER ON;  -- requerido por los índices filtrados (sqlcmd lo trae OFF)")
w("SET ANSI_NULLS ON; SET ANSI_PADDING ON; SET ANSI_WARNINGS ON; SET CONCAT_NULL_YIELDS_NULL ON; SET ARITHABORT ON;")
w("SET XACT_ABORT ON;")
w("IF DB_NAME() <> N'DB_Laboratorios_Univalle_DEMO_QA' THROW 50000, 'Destino incorrecto: solo DEMO_QA', 1;")
w("IF NOT EXISTS (SELECT 1 FROM dbo.__EFMigrationsHistory WHERE MigrationId = N'20261005144021_AddPracticeClosureRole') "
  "THROW 50001, 'Falta la migración AddPracticeClosureRole; abortar.', 1;")
w("DECLARE @Now datetime2 = SYSUTCDATETIME();")
w("DECLARE @n int;")
w(f"IF NOT EXISTS (SELECT 1 FROM dbo.ImportBatches WHERE Code = N'{batch_code}') "
  f"INSERT dbo.ImportBatches(Code, SourceType, SourceName, SourceSha256, ImportedAt, Status, Notes, ContractVersion) "
  f"VALUES(N'{batch_code}', N'Excel', N'{SOURCE_NAME}', N'{sha}', @Now, 1, "
  f"N'V12 bloque A: materias, docentes, artículos, UNI-05, vida útil, existencias y fechas de gestión. Solo DEMO_QA.', "
  f"N'base-gastronomia-v12');")
w(f"DECLARE @Batch int = (SELECT Id FROM dbo.ImportBatches WHERE Code = N'{batch_code}');")
w("CREATE TABLE #Issue(IssueCode nvarchar(100), EntityName nvarchar(100), EntityKey nvarchar(200), FieldName nvarchar(100), "
  "OriginalValue nvarchar(400), NormalizedValue nvarchar(400), Description nvarchar(1000), Severity int, SourceSheet nvarchar(128), SourceRow int);")


def issue(icode, entity, key, field, original, normalized, desc, severity, sheet_name, row):
    w(f"INSERT #Issue VALUES({q(icode)}, {q(entity)}, {q(key)}, {q(field)}, {q(original)}, {q(normalized)}, "
      f"{q(desc)}, {severity}, {q(sheet_name)}, {row if row is not None else 'NULL'});")


# 1. Unidad UNI-05
w("-- 1. Unidad de medida UNI-05 (Metro): la usa ART-REQ-00615")
w("BEGIN TRANSACTION;")
w("IF NOT EXISTS (SELECT 1 FROM dbo.MeasurementUnits WHERE Code = N'UNI-05') "
  "INSERT dbo.MeasurementUnits(Code, Name, Status, CreatedDate) VALUES(N'UNI-05', N'Metro', 0, @Now);")
w("SELECT 'unidad_uni05' k, COUNT(*) n FROM dbo.MeasurementUnits WHERE Code = N'UNI-05';")
w("COMMIT TRANSACTION;")

# 2. Materias (D3: MAYÚSCULAS)
w("-- 2. Materias nuevas (22). Nombre en MAYÚSCULAS como las 19 existentes (D3). No se toca ninguna existente.")
w("CREATE TABLE #Subj(Code nvarchar(50) PRIMARY KEY, Name nvarchar(255), Career nvarchar(50), Semester int NULL, St int, SourceRow int);")
for s in new_subjects:
    st = 0 if str(s["ESTADO"] or "").strip().lower() == "activo" else 1
    w(f"INSERT #Subj VALUES({q(s['MATERIACODIGO'])}, {q(text(s['NOMBRE']).upper())}, {q(s['CARRERACODIGO'])}, "
      f"{q(s['SEMESTRE'])}, {st}, {s['_row']});")
w("""INSERT #Issue SELECT 'UNRESOLVED_REFERENCE', 'Subject', s.Code, 'Carreracodigo', s.Career, NULL,
   N'Carrera inexistente; la materia no se carga.', 2, 'Materias', s.SourceRow
FROM #Subj s WHERE NOT EXISTS (SELECT 1 FROM dbo.Careers c WHERE c.Code = s.Career);
INSERT #Issue SELECT 'CONTENT_CONFLICT', 'Subject', s.Code, 'Nombre', s.Name, x.Code,
   N'Ya existe otra materia con el mismo nombre en la carrera; no se carga.', 2, 'Materias', s.SourceRow
FROM #Subj s JOIN dbo.Careers c ON c.Code = s.Career JOIN dbo.Subjects x ON x.CareerId = c.Id AND x.Name = s.Name AND x.Code <> s.Code;
BEGIN TRANSACTION;
INSERT dbo.Subjects(Name, Code, CareerId, Semester, Status)
SELECT s.Name, s.Code, c.Id, s.Semester, s.St
FROM #Subj s JOIN dbo.Careers c ON c.Code = s.Career
WHERE NOT EXISTS (SELECT 1 FROM dbo.Subjects x WHERE x.Code = s.Code)
  AND NOT EXISTS (SELECT 1 FROM dbo.Subjects x WHERE x.CareerId = c.Id AND x.Name = s.Name);
SELECT 'materias_insertadas' k, @@ROWCOUNT n;
COMMIT TRANSACTION;""")

# 3. Docentes nuevos (patrón V3: People + Interns + rol + alias)
w("-- 3. Docentes nuevos (6): People + Interns + rol Docente + alias. MAYÚSCULAS como los 15 existentes (D3).")
w("BEGIN TRANSACTION;")
for a in new_actors:
    c = a["ACTORCODIGO"]
    status = 0 if str(a["ESTADO"] or "Activo").strip().lower() == "activo" else 1
    pid = f"(SELECT Id FROM dbo.People WHERE ActorCode = N'{c}')"
    w(f"IF NOT EXISTS (SELECT 1 FROM dbo.People WHERE ActorCode = N'{c}') "
      f"INSERT dbo.People(Status, Category, Email, PhoneNumber, CreatedDate, ImportBatchId, ActorCode, Title) "
      f"VALUES({status}, 2, {q(text(a['EMAIL']))}, {q(text(a['TELEFONO']))}, @Now, @Batch, N'{c}', {q(a['_title'])});")
    w(f"IF NOT EXISTS (SELECT 1 FROM dbo.Interns WHERE Id = {pid}) AND NOT EXISTS (SELECT 1 FROM dbo.Externs WHERE Id = {pid}) "
      f"INSERT dbo.Interns(Id, Name, InternStatus) VALUES({pid}, {q(a['_name'])}, {status});")
    for role in sorted(roles_by_actor[c]):
        w(f"IF NOT EXISTS (SELECT 1 FROM dbo.PersonRoleAssignments WHERE PersonId = {pid} AND Role = {role} AND IsActive = 1) "
          f"INSERT dbo.PersonRoleAssignments(PersonId, Role, IsActive, ValidFrom) VALUES({pid}, {role}, 1, @Now);")
    aliases = [(a["_name"], 1)]
    if a["_title"]:
        aliases.append((f"{a['_title'].rstrip('.')} {a['_name']}", 0))
    raw = text(a["NOMBRERAZONSOCIAL"])
    if " ".join(raw.upper().split()) not in {x for x, _ in aliases}:
        aliases.append((raw, 0))
    for alias, preferred in aliases:
        normalized = " ".join(str(alias).upper().split())
        w(f"IF NOT EXISTS (SELECT 1 FROM dbo.PersonAliases WHERE PersonId = {pid} AND NormalizedAlias = {q(normalized)}) "
          f"INSERT dbo.PersonAliases(PersonId, Alias, NormalizedAlias, IsPreferred, Source) "
          f"VALUES({pid}, {q(alias)}, {q(normalized)}, {preferred}, N'{SOURCE_NAME}');")
    if str(a["VINCULACION"] or "").strip().lower() == "por confirmar":
        issue("ACTOR_LINK_PENDING", "Person", c, "Vinculacion", text(a["VINCULACION"]), "Intern",
              "La fuente marca la vinculación como 'Por confirmar': se carga como interno.", 1, "Actores", a["_row"])
    if a["_title"] and not text(a["TITULO"]):
        issue("VALUE_NORMALIZED", "Person", c, "Titulo", text(a["NOMBRERAZONSOCIAL"]), a["_title"],
              "El título venía pegado al nombre; se separó en People.Title.", 0, "Actores", a["_row"])
w("COMMIT TRANSACTION;")
w("SELECT 'docentes_v12' k, COUNT(*) n FROM dbo.People WHERE ActorCode BETWEEN N'ACT-0052' AND N'ACT-0057';")
issue("POSSIBLE_PERSON_DUPLICATE", "Person", "ACT-0052", "Nombrerazonsocial", "Ing.zabala", "ACT-0037",
      "Probable duplicado de ACT-0037 (ALBERTO ZABALA, ING.). No se fusiona: confirmar con la clienta.", 1, "Actores",
      next(a["_row"] for a in new_actors if a["ACTORCODIGO"] == "ACT-0052"))

# 4. Artículos nuevos
w("-- 4. Artículos nuevos (372). Nombre como viene (D3); el duplicado lleva el código como sufijo (D4).")
w("CREATE TABLE #Art(Code nvarchar(30) PRIMARY KEY, Name nvarchar(200), Cat nvarchar(50), Unit nvarchar(50), Review int, SourceRow int);")
for r in new_articles:
    review = 0 if str(r["ESTADO"] or "").strip().lower() == "activo" else 9
    w(f"INSERT #Art VALUES({q(r['ARTICULOCODIGO'])}, {q(r['_name'])}, {q(r['CATEGORIACODIGO'])}, "
      f"{q(r['UNIDADMEDIDACODIGO'])}, {review}, {r['_row']});")
    if review == 9:
        issue("ARTICLE_PENDING_CONFIRMATION", "Article", r["ARTICULOCODIGO"], "Estado", text(r["ESTADO"]), "ReviewStatus=9",
              "La fuente marca el artículo como 'Por confirmar': se importa Activo con ReviewStatus = Por confirmar.",
              1, "Articulos", r["_row"])
for r in dups:
    issue("DUPLICATE_NAME_IN_SOURCE", "Article", r["ARTICULOCODIGO"], "Nombre", text(r["NOMBRE"]), r["_name"],
          f"Mismo nombre que {r['_dup_of']}; se agrega el código al nombre para respetar UX_Articles_Name_Active (D4).",
          1, "Articulos", r["_row"])
w("""INSERT #Issue SELECT 'UNRESOLVED_REFERENCE', 'Article', a.Code, 'Categoria/Unidad', CONCAT(a.Cat, N' / ', a.Unit), NULL,
   N'Categoría o unidad inexistente; el artículo no se carga.', 2, 'Articulos', a.SourceRow
FROM #Art a WHERE NOT EXISTS (SELECT 1 FROM dbo.ArticleCategories c WHERE c.Code = a.Cat)
             OR NOT EXISTS (SELECT 1 FROM dbo.MeasurementUnits u WHERE u.Code = a.Unit);
INSERT #Issue SELECT 'CONTENT_CONFLICT', 'Article', a.Code, 'Nombre', a.Name, x.Code,
   N'Ya existe un artículo activo con el mismo nombre; no se carga.', 2, 'Articulos', a.SourceRow
FROM #Art a JOIN dbo.Articles x ON x.Name = a.Name AND x.Status <> 2 AND x.Code <> a.Code;
SET @n = 1;
WHILE @n > 0
BEGIN
    BEGIN TRANSACTION;
    INSERT dbo.Articles(Code, Name, Category, UnitOfMeasure, Status, ArticleCategoryId, MeasurementUnitId, ReviewStatus, ImportBatchId, CreatedDate)
    SELECT TOP (100) a.Code, a.Name, c.Name, u.Name, 0, c.Id, u.Id, a.Review, @Batch, @Now
    FROM #Art a JOIN dbo.ArticleCategories c ON c.Code = a.Cat JOIN dbo.MeasurementUnits u ON u.Code = a.Unit
    WHERE NOT EXISTS (SELECT 1 FROM dbo.Articles x WHERE x.Code = a.Code)
      AND NOT EXISTS (SELECT 1 FROM dbo.Articles x WHERE x.Name = a.Name AND x.Status <> 2);
    SET @n = @@ROWCOUNT;
    COMMIT TRANSACTION;
END;
SELECT 'articulos_v12' k, COUNT(*) n FROM dbo.Articles a JOIN #Art t ON t.Code = a.Code;""")

# 5. Vida útil
w("-- 5. Vida útil del catálogo de equipos (121)")
w("CREATE TABLE #Life(Code nvarchar(30) PRIMARY KEY, Years int, SourceRow int);")
for c in catalog:
    w(f"INSERT #Life VALUES({q(c['CATALOGOCODIGO'])}, {int(c['VIDAUTILANIOS'])}, {c['_row']});")
w("""INSERT #Issue SELECT 'UNRESOLVED_REFERENCE', 'Equipment', l.Code, 'Catalogocodigo', l.Code, NULL,
   N'Código de catálogo inexistente; no se actualiza la vida útil.', 2, 'CatalogoEquipos', l.SourceRow
FROM #Life l WHERE NOT EXISTS (SELECT 1 FROM dbo.Equipments e WHERE e.CatalogCode = l.Code);
INSERT #Issue SELECT 'CONTENT_CONFLICT', 'Equipment', l.Code, 'Vidautilanios', CAST(e.UsefulLifeYears AS nvarchar(20)), CAST(l.Years AS nvarchar(20)),
   N'La BD tenía otra vida útil; se reemplaza por la de la fuente.', 1, 'CatalogoEquipos', l.SourceRow
FROM #Life l JOIN dbo.Equipments e ON e.CatalogCode = l.Code
WHERE e.UsefulLifeYears IS NOT NULL AND e.UsefulLifeYears <> l.Years;
BEGIN TRANSACTION;
UPDATE e SET UsefulLifeYears = l.Years, LastModifiedDate = @Now
FROM dbo.Equipments e JOIN #Life l ON l.Code = e.CatalogCode
WHERE e.UsefulLifeYears IS NULL OR e.UsefulLifeYears <> l.Years;
SELECT 'vida_util_actualizada' k, @@ROWCOUNT n;
COMMIT TRANSACTION;""")

# 6. Existencias (D5: agregar y actualizar, no borrar)
w("-- 6. Existencias (D5): insertar las nuevas, actualizar las distintas, conservar las que solo están en la BD")
w("CREATE TABLE #Stk(Lab nvarchar(50), Art nvarchar(30), Qty decimal(18,2), Upd date NULL, SourceRow int, PRIMARY KEY(Lab, Art));")
for s in stocks:
    w(f"INSERT #Stk VALUES({q(s['LABORATORIOCODIGO'])}, {q(s['ARTICULOCODIGO'])}, {q(s['CANTIDADDISPONIBLE'])}, "
      f"{q(s['FECHAACTUALIZACION'])}, {s['_row']});")
w("""INSERT #Issue SELECT 'UNRESOLVED_REFERENCE', 'ArticleStock', CONCAT(s.Lab, N'/', s.Art), 'Laboratorio/Articulo', NULL, NULL,
   N'Laboratorio o artículo inexistente o dado de baja; no se carga.', 2, 'ExistenciasArticulos', s.SourceRow
FROM #Stk s WHERE NOT EXISTS (SELECT 1 FROM dbo.Laboratories l WHERE l.Code = s.Lab)
             OR NOT EXISTS (SELECT 1 FROM dbo.Articles a WHERE a.Code = s.Art AND a.Status <> 2);
-- Cantidad distinta: se actualiza y se deja el valor anterior en la incidencia.
INSERT #Issue SELECT 'CONTENT_CONFLICT', 'ArticleStock', CONCAT(s.Lab, N'/', s.Art), 'Cantidaddisponible',
   CAST(x.AvailableQuantity AS nvarchar(40)), CAST(s.Qty AS nvarchar(40)),
   N'La V12 trae otra cantidad: se actualiza; el valor anterior queda en OriginalValue (D5).', 1, 'ExistenciasArticulos', s.SourceRow
FROM #Stk s JOIN dbo.Laboratories l ON l.Code = s.Lab JOIN dbo.Articles a ON a.Code = s.Art
JOIN dbo.ArticleStocks x ON x.LaboratoryId = l.Id AND x.ArticleId = a.Id
WHERE x.AvailableQuantity <> s.Qty
  AND NOT EXISTS (SELECT 1 FROM dbo.DataQualityIssues d WHERE d.ImportBatchId = @Batch AND d.IssueCode = N'CONTENT_CONFLICT'
                  AND d.EntityKey = CONCAT(s.Lab, N'/', s.Art) AND d.FieldName = N'Cantidaddisponible');
-- Existencias de la BD que la V12 ya no trae: se conservan.
INSERT #Issue SELECT 'STOCK_NOT_IN_SOURCE', 'ArticleStock', CONCAT(l.Code, N'/', a.Code), 'Cantidaddisponible',
   CAST(x.AvailableQuantity AS nvarchar(40)), NULL,
   N'La existencia está en la BD pero no en la V12; se conserva sin cambios (D5). Confirmar con la clienta.', 1, 'ExistenciasArticulos', NULL
FROM dbo.ArticleStocks x JOIN dbo.Laboratories l ON l.Id = x.LaboratoryId JOIN dbo.Articles a ON a.Id = x.ArticleId
WHERE NOT EXISTS (SELECT 1 FROM #Stk s WHERE s.Lab = l.Code AND s.Art = a.Code);
BEGIN TRANSACTION;
UPDATE x SET AvailableQuantity = s.Qty, UpdatedOn = COALESCE(s.Upd, x.UpdatedOn), LastModifiedDate = @Now
FROM dbo.ArticleStocks x JOIN dbo.Laboratories l ON l.Id = x.LaboratoryId JOIN dbo.Articles a ON a.Id = x.ArticleId
JOIN #Stk s ON s.Lab = l.Code AND s.Art = a.Code
WHERE x.AvailableQuantity <> s.Qty;
SELECT 'existencias_actualizadas' k, @@ROWCOUNT n;
UPDATE x SET UpdatedOn = s.Upd, LastModifiedDate = @Now
FROM dbo.ArticleStocks x JOIN dbo.Laboratories l ON l.Id = x.LaboratoryId JOIN dbo.Articles a ON a.Id = x.ArticleId
JOIN #Stk s ON s.Lab = l.Code AND s.Art = a.Code
WHERE x.UpdatedOn IS NULL AND s.Upd IS NOT NULL;
SELECT 'existencias_fecha_completada' k, @@ROWCOUNT n;
COMMIT TRANSACTION;
SET @n = 1;
WHILE @n > 0
BEGIN
    BEGIN TRANSACTION;
    INSERT dbo.ArticleStocks(LaboratoryId, ArticleId, AvailableQuantity, UpdatedOn, CreatedDate)
    SELECT TOP (100) l.Id, a.Id, s.Qty, s.Upd, @Now
    FROM #Stk s JOIN dbo.Laboratories l ON l.Code = s.Lab JOIN dbo.Articles a ON a.Code = s.Art AND a.Status <> 2
    WHERE NOT EXISTS (SELECT 1 FROM dbo.ArticleStocks x WHERE x.LaboratoryId = l.Id AND x.ArticleId = a.Id);
    SET @n = @@ROWCOUNT;
    COMMIT TRANSACTION;
END;
SELECT 'existencias_total' k, COUNT(*) n FROM dbo.ArticleStocks;""")

# 7. Fechas de gestión (D7: solo donde son NULL)
w("-- 7. Fechas de gestión (D7): solo se completan las que están en NULL en la BD")
w("CREATE TABLE #Mgmt(Code nvarchar(50) PRIMARY KEY, StartDate date, EndDate date, SourceRow int);")
for g in B["Gestiones"]:
    if g["FECHAINICIO"] and g["FECHAFIN"]:
        w(f"INSERT #Mgmt VALUES({q(g['GESTIONCODIGO'])}, {q(g['FECHAINICIO'])}, {q(g['FECHAFIN'])}, {g['_row']});")
w("""INSERT #Issue SELECT 'VALUE_FILLED_FROM_SOURCE', 'Management', m.Code, 'Fechainicio/Fechafin', NULL,
   CONCAT(CONVERT(nvarchar(10), g.StartDate, 23), N' / ', CONVERT(nvarchar(10), g.EndDate, 23)),
   N'La gestión no tenía fechas; se completan con las de la fuente (D7). Son fechas calendario, no de inicio de clases.', 0, 'Gestiones', g.SourceRow
FROM dbo.Managements m JOIN #Mgmt g ON g.Code = m.Code
WHERE m.StartDate IS NULL AND m.PlannedEndDate IS NULL;
BEGIN TRANSACTION;
UPDATE m SET StartDate = g.StartDate, PlannedEndDate = g.EndDate, LastModifiedDate = @Now
FROM dbo.Managements m JOIN #Mgmt g ON g.Code = m.Code
WHERE m.StartDate IS NULL AND m.PlannedEndDate IS NULL;
SELECT 'gestiones_con_fecha' k, @@ROWCOUNT n;
COMMIT TRANSACTION;""")

# 8. Incidencias (idempotentes)
w("""-- 8. Incidencias (idempotentes por ImportBatchId + IssueCode + EntityKey + FieldName)
SET @n = 1;
WHILE @n > 0
BEGIN
    BEGIN TRANSACTION;
    INSERT dbo.DataQualityIssues(ImportBatchId, IssueCode, EntityName, EntityKey, FieldName, OriginalValue, NormalizedValue,
        Description, Severity, Status, CreatedDate, SourceSheet, SourceRowNumber)
    SELECT TOP (100) @Batch, i.IssueCode, i.EntityName, i.EntityKey, i.FieldName, i.OriginalValue, i.NormalizedValue,
        i.Description, i.Severity, 0, @Now, i.SourceSheet, i.SourceRow
    FROM (SELECT DISTINCT * FROM #Issue) i
    WHERE NOT EXISTS (SELECT 1 FROM dbo.DataQualityIssues d WHERE d.ImportBatchId = @Batch AND d.IssueCode = i.IssueCode
                      AND ISNULL(d.EntityKey, N'') = ISNULL(i.EntityKey, N'')
                      AND ISNULL(d.FieldName, N'') = ISNULL(i.FieldName, N''));
    SET @n = @@ROWCOUNT;
    COMMIT TRANSACTION;
END;
SELECT IssueCode, COUNT(*) n FROM dbo.DataQualityIssues WHERE ImportBatchId = @Batch GROUP BY IssueCode ORDER BY IssueCode;""")

# 9. Verificación
w("""-- 9. Verificación
SELECT 'materias_activas' k, COUNT(*) n FROM dbo.Subjects WHERE Status = 0
UNION ALL SELECT 'actores_act', COUNT(*) FROM dbo.People WHERE ActorCode LIKE N'ACT-%'
UNION ALL SELECT 'docentes_act_rol8', COUNT(*) FROM dbo.PersonRoleAssignments r JOIN dbo.People p ON p.Id = r.PersonId
          WHERE r.Role = 8 AND r.IsActive = 1 AND p.ActorCode LIKE N'ACT-%'
UNION ALL SELECT 'articulos_activos_sin_demo', COUNT(*) FROM dbo.Articles WHERE Status <> 2 AND Code NOT LIKE N'ART-000%'
UNION ALL SELECT 'equipos_con_vida_util', COUNT(*) FROM dbo.Equipments WHERE UsefulLifeYears IS NOT NULL
UNION ALL SELECT 'existencias', COUNT(*) FROM dbo.ArticleStocks
UNION ALL SELECT 'gestiones_sin_fecha', COUNT(*) FROM dbo.Managements WHERE StartDate IS NULL;""")

os.makedirs(os.path.dirname(OUT), exist_ok=True)
io.open(OUT, "w", encoding="utf-8").write("\n".join(sql) + "\n")
print(OUT, f"| lote={batch_code}", flush=True)
