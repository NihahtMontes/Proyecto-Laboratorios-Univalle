import os
from openpyxl import load_workbook

folder = "wwwroot/templates"
for filename in os.listdir(folder):
    if filename.endswith(".xlsx") and not filename.startswith("solicitud"):
        print(f"\n=====================================")
        print(f"ARCHIVO: {filename}")
        print(f"=====================================")
        filepath = os.path.join(folder, filename)
        try:
            wb = load_workbook(filename=filepath, read_only=True, data_only=True)
            ws = wb.active
            for row in ws.iter_rows(values_only=False, max_row=60, max_col=20):
                for cell in row:
                    if cell.value is not None and str(cell.value).strip() != "":
                        val = str(cell.value).strip().replace('\n', ' ')
                        if len(val) > 80:
                            val = val[:80] + "..."
                        print(f"Celda {cell.coordinate}: {val}")
        except Exception as e:
            print(f"Error reading {filename}: {e}")
